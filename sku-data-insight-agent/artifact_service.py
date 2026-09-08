from __future__ import annotations

import hashlib
import json
import mimetypes
import re
import shutil
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from memory_store import MemoryStore


SUPPORTED = {".pdf": "pdf", ".pptx": "pptx", ".docx": "docx", ".csv": "csv", ".json": "json", ".txt": "txt", ".md": "txt", ".xlsx": "xlsx"}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


class ArtifactService:
    def __init__(self, root: Path, store: MemoryStore) -> None:
        self.root, self.store = root, store
        self.files = root / "artifacts"
        self.files.mkdir(parents=True, exist_ok=True)

    def ingest(self, source: Path, workspace_id: str, user_id: str = "admin") -> dict[str, Any]:
        suffix = source.suffix.lower()
        if suffix not in SUPPORTED:
            raise ValueError(f"Unsupported artifact type: {suffix}")
        digest = sha256_file(source)
        existing = next((item for item in self.store.list_artifacts(workspace_id, user_id) if item["source_sha256"] == digest), None)
        if existing:
            return existing
        artifact_id = f"artifact-{uuid.uuid4().hex}"
        folder = self.files / artifact_id
        folder.mkdir(parents=True)
        destination = folder / source.name
        shutil.copy2(source, destination)
        manifest = self._extract(destination, suffix, folder)
        record = {"artifact_id": artifact_id, "workspace_id": workspace_id, "user_id": user_id, "file_name": source.name, "media_type": SUPPORTED[suffix], "source_sha256": digest, "storage_key": str(destination), "extraction_status": manifest.get("status", "completed"), "manifest": manifest, "created_at": datetime.now(timezone.utc).isoformat()}
        return self.store.add_artifact(record)

    def path_for(self, artifact: dict[str, Any]) -> Path:
        path = Path(artifact["storage_key"]).resolve()
        root = self.files.resolve()
        if root not in path.parents or not path.exists():
            raise FileNotFoundError("Artifact is not available")
        return path

    def _extract(self, path: Path, suffix: str, folder: Path) -> dict[str, Any]:
        try:
            if suffix == ".pdf":
                from pypdf import PdfReader
                pages = []
                for index, page in enumerate(PdfReader(str(path)).pages, 1):
                    pages.append({"locator": f"page:{index}", "text": (page.extract_text() or "")[:12000]})
                return {"status": "completed", "kind": "pdf", "page_count": len(pages), "pages": pages}
            if suffix == ".pptx":
                from pptx import Presentation
                slides = []
                for index, slide in enumerate(Presentation(str(path)).slides, 1):
                    text = "\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text).strip()
                    slides.append({"locator": f"slide:{index}", "text": text[:12000]})
                return {"status": "completed", "kind": "pptx", "slide_count": len(slides), "slides": slides}
            if suffix == ".docx":
                from docx import Document
                paragraphs = [item.text for item in Document(str(path)).paragraphs if item.text.strip()]
                return {"status": "completed", "kind": "docx", "paragraphs": paragraphs[:2000]}
            if suffix in {".txt", ".md", ".csv", ".json"}:
                text = path.read_text(encoding="utf-8-sig", errors="replace")[:200000]
                return {"status": "completed", "kind": SUPPORTED[suffix], "text": text}
            return {"status": "registered", "kind": SUPPORTED[suffix], "note": "Strict Excel ingestion remains owned by WorkbookAdapter"}
        except ImportError as exc:
            return {"status": "blocked", "kind": SUPPORTED[suffix], "error": f"Optional parser dependency missing: {exc.name}"}
        except Exception as exc:
            return {"status": "blocked", "kind": SUPPORTED[suffix], "error": str(exc)[:500]}

    @staticmethod
    def evidence(artifact: dict[str, Any], limit: int = 12) -> list[dict[str, Any]]:
        manifest = artifact.get("manifest", {})
        result: list[dict[str, Any]] = []
        for page in manifest.get("pages", [])[:limit]:
            result.append({"artifact_id": artifact["artifact_id"], "locator_type": "page", "locator": page["locator"], "excerpt": page.get("text", "")[:4000], "content_sha256": hashlib.sha256(page.get("text", "").encode("utf-8")).hexdigest()})
        for slide in manifest.get("slides", [])[:limit]:
            result.append({"artifact_id": artifact["artifact_id"], "locator_type": "slide", "locator": slide["locator"], "excerpt": slide.get("text", "")[:4000], "content_sha256": hashlib.sha256(slide.get("text", "").encode("utf-8")).hexdigest()})
        if manifest.get("text"):
            result.append({"artifact_id": artifact["artifact_id"], "locator_type": "document", "locator": "text:1", "excerpt": manifest["text"][:4000], "content_sha256": hashlib.sha256(manifest["text"].encode("utf-8")).hexdigest()})
        if manifest.get("paragraphs"):
            result.append({"artifact_id": artifact["artifact_id"], "locator_type": "paragraph", "locator": "paragraph:1-" + str(min(len(manifest["paragraphs"]), 20)), "excerpt": "\n".join(manifest["paragraphs"][:20])[:4000], "content_sha256": hashlib.sha256("\n".join(manifest["paragraphs"][:20]).encode("utf-8")).hexdigest()})
        return result
