from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path
from typing import Any, Literal

from openpyxl.utils.cell import coordinate_to_tuple
from pydantic import BaseModel, Field, field_validator, model_validator


def normalize_label(value: Any) -> str:
    return "".join(character for character in str(value or "").upper() if character.isalnum())


class ManifestCell(BaseModel):
    sheet: str
    cell: str
    row: int = Field(ge=1)
    column: int = Field(ge=1)
    raw_text: str = Field(min_length=1, max_length=200)
    style_id: int = Field(ge=0)
    bold: bool = False


class WorkbookStructureManifest(BaseModel):
    manifest_version: Literal["1.0"] = "1.0"
    mapping_sheet: str
    sheet_names: list[str]
    max_row: int = Field(ge=1)
    max_column: int = Field(ge=1)
    merged_ranges: list[str]
    cells: list[ManifestCell]
    layout_fingerprint: str = Field(pattern=r"^[a-f0-9]{64}$")


class SemanticLabel(BaseModel):
    cell: str
    raw_text: str = Field(min_length=1, max_length=200)
    role: Literal["CHANNEL", "PRIMARY_SKU_HEADER", "COMPETITOR_BRAND_HEADER"]
    canonical_value: str = Field(min_length=1, max_length=80)
    confidence: float = Field(ge=0, le=1)
    novel: bool = False

    @field_validator("canonical_value")
    @classmethod
    def normalize_canonical(cls, value: str) -> str:
        return value.strip().upper()


class CompetitorHeader(BaseModel):
    brand: str = Field(min_length=1, max_length=80)
    label: SemanticLabel

    @field_validator("brand")
    @classmethod
    def uppercase_brand(cls, value: str) -> str:
        return value.strip().upper()


class ChannelGroupPlan(BaseModel):
    channel: Literal["JD", "ALI", "OFFLINE"]
    raw_sheet: str
    channel_label: SemanticLabel
    header_row: int = Field(ge=1)
    start_column: int = Field(ge=1)
    end_column: int = Field(ge=1)
    data_start_row: int = Field(ge=1)
    data_end_row: int = Field(ge=1)
    primary_header: SemanticLabel
    competitor_headers: list[CompetitorHeader]
    continuation_rule: Literal["MERGED_OR_BLANK_PRIMARY_UNTIL_NEXT"] = "MERGED_OR_BLANK_PRIMARY_UNTIL_NEXT"

    @model_validator(mode="after")
    def valid_bounds(self) -> "ChannelGroupPlan":
        if self.start_column > self.end_column or self.data_start_row > self.data_end_row:
            raise ValueError("Invalid channel group bounds")
        if self.header_row >= self.data_start_row:
            raise ValueError("Data rows must start after the header")
        return self


class WorkbookSemanticPlan(BaseModel):
    plan_version: Literal["1.0"] = "1.0"
    layout_fingerprint: str = Field(pattern=r"^[a-f0-9]{64}$")
    source: Literal["deterministic", "agnes", "approved_cache"]
    channel_groups: list[ChannelGroupPlan] = Field(min_length=1)
    unknown_labels: list[SemanticLabel] = Field(default_factory=list)
    review_required: bool = False
    review_reasons: list[str] = Field(default_factory=list)
    human_approved: bool = False
    recognizer_model: str | None = None
    prompt_version: str = "semantic-plan-v1"

    @model_validator(mode="after")
    def unique_channels(self) -> "WorkbookSemanticPlan":
        channels = [group.channel for group in self.channel_groups]
        if len(channels) != len(set(channels)):
            raise ValueError("Semantic plan contains duplicate channel groups")
        return self

    def group_for(self, channel: str) -> ChannelGroupPlan:
        target = channel.upper()
        for group in self.channel_groups:
            if group.channel == target:
                return group
        raise KeyError(channel)


class WorkbookSemanticProfiler:
    def __init__(
        self,
        workbook: Any,
        mapping_sheet: str,
        channel_sheets: dict[str, str],
        channel_aliases: dict[str, list[str] | tuple[str, ...]],
        primary_header_aliases: list[str] | tuple[str, ...],
        confidence_threshold: float = 0.90,
    ) -> None:
        self.workbook = workbook
        self.mapping_sheet = mapping_sheet
        self.channel_sheets = {key.upper(): value for key, value in channel_sheets.items()}
        self.channel_aliases = {
            channel.upper(): {normalize_label(alias) for alias in aliases}
            for channel, aliases in channel_aliases.items()
        }
        self.primary_header_aliases = {normalize_label(value) for value in primary_header_aliases}
        self.confidence_threshold = confidence_threshold
        self.sheet = workbook[mapping_sheet]

    def build_manifest(self) -> WorkbookStructureManifest:
        cells: list[ManifestCell] = []
        for row in self.sheet.iter_rows():
            for cell in row:
                value = cell.value
                if value in (None, "") or not isinstance(value, str):
                    continue
                text = value.strip()
                if not text:
                    continue
                cells.append(
                    ManifestCell(
                        sheet=self.mapping_sheet,
                        cell=cell.coordinate,
                        row=cell.row,
                        column=cell.column,
                        raw_text=text[:200],
                        style_id=int(cell.style_id),
                        bold=bool(cell.font and cell.font.bold),
                    )
                )
        if len(cells) > 5000:
            raise ValueError("Mapping sheet contains too many label cells")
        fingerprint_payload = {
            "mapping_sheet": self.mapping_sheet,
            "sheet_names": list(self.workbook.sheetnames),
            "shape": [self.sheet.max_row, self.sheet.max_column],
            "merged_ranges": sorted(str(item) for item in self.sheet.merged_cells.ranges),
            "cells": [[item.cell, item.raw_text, item.style_id, item.bold] for item in cells],
        }
        fingerprint = hashlib.sha256(
            json.dumps(fingerprint_payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
        ).hexdigest()
        return WorkbookStructureManifest(
            mapping_sheet=self.mapping_sheet,
            sheet_names=list(self.workbook.sheetnames),
            max_row=self.sheet.max_row,
            max_column=self.sheet.max_column,
            merged_ranges=sorted(str(item) for item in self.sheet.merged_cells.ranges),
            cells=cells,
            layout_fingerprint=fingerprint,
        )

    def deterministic_plan(self, manifest: WorkbookStructureManifest | None = None) -> WorkbookSemanticPlan:
        manifest = manifest or self.build_manifest()
        channel_anchors: dict[str, tuple[int, int, str]] = {}
        for row in range(1, self.sheet.max_row + 1):
            for column in range(1, self.sheet.max_column + 1):
                raw = self.sheet.cell(row, column).value
                normalized = normalize_label(raw)
                for channel, aliases in self.channel_aliases.items():
                    if normalized in aliases and channel not in channel_anchors:
                        channel_anchors[channel] = (row, column, str(raw).strip())
        missing = sorted(set(self.channel_sheets) - set(channel_anchors))
        if missing:
            raise ValueError(f"Unable to recognize channel labels: {', '.join(missing)}")

        groups: list[ChannelGroupPlan] = []
        unknown: list[SemanticLabel] = []
        for channel, (channel_row, channel_column, channel_raw) in channel_anchors.items():
            primary_cell = None
            for row in range(channel_row + 1, min(self.sheet.max_row, channel_row + 4) + 1):
                for column in range(channel_column, self.sheet.max_column + 1):
                    raw = self.sheet.cell(row, column).value
                    if normalize_label(raw) in self.primary_header_aliases:
                        primary_cell = self.sheet.cell(row, column)
                        break
                if primary_cell is not None:
                    break
            if primary_cell is None:
                raise ValueError(f"Unable to recognize Philips header for {channel}")

            header_row = primary_cell.row
            start_column = primary_cell.column
            end_column = start_column
            while end_column + 1 <= self.sheet.max_column and self.sheet.cell(header_row, end_column + 1).value not in (None, ""):
                end_column += 1
            if end_column == start_column:
                raise ValueError(f"No competitor headers found for {channel}")

            later_rows = [
                anchor_row
                for other_channel, (anchor_row, anchor_column, _) in channel_anchors.items()
                if other_channel != channel
                and anchor_row > channel_row
                and start_column <= anchor_column <= end_column
            ]
            data_end_row = min(later_rows) - 1 if later_rows else self.sheet.max_row

            channel_novel = normalize_label(channel_raw) not in self.channel_aliases[channel]
            channel_label = SemanticLabel(
                cell=self.sheet.cell(channel_row, channel_column).coordinate,
                raw_text=channel_raw,
                role="CHANNEL",
                canonical_value=channel,
                confidence=1.0,
                novel=channel_novel,
            )
            primary_raw = str(primary_cell.value).strip()
            primary_novel = normalize_label(primary_raw) not in self.primary_header_aliases
            primary_label = SemanticLabel(
                cell=primary_cell.coordinate,
                raw_text=primary_raw,
                role="PRIMARY_SKU_HEADER",
                canonical_value="PHILIPS",
                confidence=1.0,
                novel=primary_novel,
            )
            competitor_headers: list[CompetitorHeader] = []
            for column in range(start_column + 1, end_column + 1):
                cell = self.sheet.cell(header_row, column)
                raw = str(cell.value).strip()
                label = SemanticLabel(
                    cell=cell.coordinate,
                    raw_text=raw,
                    role="COMPETITOR_BRAND_HEADER",
                    canonical_value=raw.upper(),
                    confidence=1.0,
                    novel=False,
                )
                competitor_headers.append(CompetitorHeader(brand=raw, label=label))
            for label in (channel_label, primary_label):
                if label.novel:
                    unknown.append(label)
            groups.append(
                ChannelGroupPlan(
                    channel=channel,
                    raw_sheet=self.channel_sheets[channel],
                    channel_label=channel_label,
                    header_row=header_row,
                    start_column=start_column,
                    end_column=end_column,
                    data_start_row=header_row + 1,
                    data_end_row=data_end_row,
                    primary_header=primary_label,
                    competitor_headers=competitor_headers,
                )
            )
        groups.sort(key=lambda item: (coordinate_to_tuple(item.channel_label.cell), item.channel))
        return WorkbookSemanticPlan(
            layout_fingerprint=manifest.layout_fingerprint,
            source="deterministic",
            channel_groups=groups,
            unknown_labels=unknown,
            review_required=bool(unknown),
            review_reasons=["Novel channel or primary header labels require confirmation"] if unknown else [],
        )

    def validate_plan(self, plan: WorkbookSemanticPlan, manifest: WorkbookStructureManifest) -> WorkbookSemanticPlan:
        if plan.layout_fingerprint != manifest.layout_fingerprint:
            raise ValueError("Semantic plan fingerprint does not match the uploaded workbook")
        manifest_cells = {item.cell: item for item in manifest.cells}
        expected_channels = set(self.channel_sheets)
        actual_channels = {item.channel for item in plan.channel_groups}
        if actual_channels != expected_channels:
            raise ValueError(f"Semantic plan channels must be {sorted(expected_channels)}")

        boxes: list[tuple[str, int, int, int, int]] = []
        computed_unknown: list[SemanticLabel] = []
        review_reasons = list(plan.review_reasons)
        for group in plan.channel_groups:
            if group.raw_sheet != self.channel_sheets[group.channel]:
                raise ValueError(f"Raw sheet mismatch for {group.channel}")
            if group.data_end_row > manifest.max_row or group.end_column > manifest.max_column:
                raise ValueError(f"Channel sandbox exceeds the mapping sheet for {group.channel}")
            if group.channel_label.role != "CHANNEL" or group.channel_label.canonical_value != group.channel:
                raise ValueError(f"Channel label contract mismatch for {group.channel}")
            if group.primary_header.canonical_value != "PHILIPS":
                raise ValueError(f"Primary header must canonicalize to PHILIPS for {group.channel}")
            labels = [group.channel_label, group.primary_header, *[item.label for item in group.competitor_headers]]
            for label in labels:
                source = manifest_cells.get(label.cell)
                if source is None or source.raw_text != label.raw_text:
                    raise ValueError(f"Semantic label source mismatch at {label.cell}")
                row, column = coordinate_to_tuple(label.cell)
                if label.role != "CHANNEL" and (row != group.header_row or not group.start_column <= column <= group.end_column):
                    raise ValueError(f"Header label {label.cell} is outside the channel header band")
                if label.novel or label.confidence < self.confidence_threshold:
                    computed_unknown.append(label.model_copy(update={"novel": True}))
            if normalize_label(group.channel_label.raw_text) not in self.channel_aliases[group.channel]:
                computed_unknown.append(group.channel_label.model_copy(update={"novel": True}))
            if normalize_label(group.primary_header.raw_text) not in self.primary_header_aliases:
                computed_unknown.append(group.primary_header.model_copy(update={"novel": True}))
            if group.primary_header.role != "PRIMARY_SKU_HEADER" or group.primary_header.cell != self.sheet.cell(group.header_row, group.start_column).coordinate:
                raise ValueError(f"Primary header is not the first column for {group.channel}")
            expected_competitor_columns = list(range(group.start_column + 1, group.end_column + 1))
            actual_competitor_columns = [coordinate_to_tuple(item.label.cell)[1] for item in group.competitor_headers]
            if actual_competitor_columns != expected_competitor_columns:
                raise ValueError(f"Competitor header columns are incomplete for {group.channel}")
            if len({item.brand for item in group.competitor_headers}) != len(group.competitor_headers):
                raise ValueError(f"Duplicate competitor brands in {group.channel}")
            for item in group.competitor_headers:
                if item.label.role != "COMPETITOR_BRAND_HEADER" or normalize_label(item.label.raw_text) != normalize_label(item.brand):
                    raise ValueError(f"Competitor header contract mismatch at {item.label.cell}")
            if any(self.sheet.cell(group.header_row, column).value in (None, "") for column in range(group.start_column, group.end_column + 1)):
                raise ValueError(f"Header band contains a gap for {group.channel}")
            if group.end_column < manifest.max_column and self.sheet.cell(group.header_row, group.end_column + 1).value not in (None, ""):
                raise ValueError(f"Semantic plan truncates a competitor header for {group.channel}")
            boxes.append((group.channel, group.data_start_row, group.data_end_row, group.start_column, group.end_column))

        for index, left in enumerate(boxes):
            for right in boxes[index + 1 :]:
                rows_overlap = max(left[1], right[1]) <= min(left[2], right[2])
                columns_overlap = max(left[3], right[3]) <= min(left[4], right[4])
                if rows_overlap and columns_overlap:
                    raise ValueError(f"Channel sandboxes overlap: {left[0]} and {right[0]}")

        all_unknown = {item.cell: item for item in [*plan.unknown_labels, *computed_unknown]}
        if all_unknown and "Novel or low-confidence labels require confirmation" not in review_reasons:
            review_reasons.append("Novel or low-confidence labels require confirmation")
        return plan.model_copy(
            update={
                "unknown_labels": list(all_unknown.values()),
                "review_required": bool(all_unknown),
                "review_reasons": review_reasons,
            }
        )


class SemanticPlanCache:
    def __init__(self, root: Path) -> None:
        self.root = root
        self.root.mkdir(parents=True, exist_ok=True)

    def path_for(self, fingerprint: str) -> Path:
        if not re.fullmatch(r"[a-f0-9]{64}", fingerprint):
            raise ValueError("Invalid semantic plan fingerprint")
        return self.root / f"{fingerprint}.json"

    def read(self, fingerprint: str, *, prompt_version: str, recognizer_model: str | None) -> WorkbookSemanticPlan | None:
        path = self.path_for(fingerprint)
        if not path.exists():
            return None
        try:
            plan = WorkbookSemanticPlan.model_validate_json(path.read_text(encoding="utf-8"))
        except Exception:
            return None
        if plan.prompt_version != prompt_version or plan.recognizer_model != recognizer_model:
            return None
        return plan.model_copy(update={"source": "approved_cache" if plan.human_approved else plan.source})

    def write(self, plan: WorkbookSemanticPlan) -> None:
        path = self.path_for(plan.layout_fingerprint)
        temporary = path.with_suffix(".tmp")
        temporary.write_text(plan.model_dump_json(indent=2), encoding="utf-8")
        temporary.replace(path)
