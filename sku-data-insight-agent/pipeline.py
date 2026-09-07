from __future__ import annotations

import hashlib
import json
import math
import os
import re
import shutil
import statistics
import subprocess
import threading
import time
import uuid
import zipfile
from dataclasses import asdict, dataclass, replace
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable, Iterable, Literal

import duckdb
import httpx
import openpyxl
import pandas as pd
import pandera.pandas as pa
from pandera import Check
from pydantic import BaseModel, ConfigDict, Field, field_validator, model_validator
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

from adapters import MappingItem, MappingResolution, MonthlyMetric, WorkbookAdapter, WorkbookContractError, normalize, serialize_mappings, serialize_resolutions, write_json
from semantic import SemanticPlanCache, WorkbookSemanticPlan, WorkbookStructureManifest


CHANNEL_NAME_ALIASES = {"JD": ("JD", "京东"), "ALI": ("ALI", "阿里", "天猫", "淘宝"), "OFFLINE": ("OFFLINE", "线下", "门店")}
RENDER_PAYLOAD_VERSION = str(json.loads(Path(__file__).with_name("render_contract.json").read_text(encoding="utf-8"))["current_payload_version"])
PHILIPS_SKU_PATTERN = re.compile(r"(?<![A-Z0-9])(?:[A-Z]{1,5}\d{2,6})(?:/\d{2})?", re.I)


def prompt_sku_matches(prompt: str) -> list[re.Match[str]]:
    matches = list(PHILIPS_SKU_PATTERN.finditer(prompt))
    if len(matches) > 1:
        specific = [match for match in matches if not re.fullmatch(r"[SI]\d000", match.group(0), re.I)]
        return specific or matches
    return matches


class AnalysisTarget(BaseModel):
    model_config = ConfigDict(extra="forbid")

    channel: Literal["JD", "ALI", "OFFLINE"]
    primary_sku: str = Field(min_length=2, max_length=80)

    @field_validator("primary_sku")
    @classmethod
    def normalize_sku(cls, value: str) -> str:
        normalized = value.strip().upper()
        normalized = re.sub(r"^(?:PHILIPS|飞利浦)[\s:/_-]+", "", normalized, flags=re.I)
        normalized = re.sub(r"\b([A-Z]{1,5}\d{2,6})-(\d{2})\b", r"\1/\2", normalized)
        tokens = [match.group(0).upper() for match in prompt_sku_matches(normalized)]
        if len(tokens) > 1:
            return " / ".join(dict.fromkeys(tokens))
        if len(tokens) == 1:
            return tokens[0]
        return normalized


class AnalysisIntent(BaseModel):
    model_config = ConfigDict(extra="forbid")

    analysis_targets: list[AnalysisTarget] = Field(min_length=1)
    months: int = Field(default=12, ge=3, le=24)
    metrics: list[Literal["ASP", "Unit"]] = Field(default_factory=lambda: ["ASP", "Unit"])
    competitor_policy: Literal["one_model_per_brand"] = "one_model_per_brand"
    mapping_policy: Literal["key_sku_list_first"] = "key_sku_list_first"
    output_template: Literal["sku_insight_v20"] = "sku_insight_v20"

    @model_validator(mode="before")
    @classmethod
    def accept_legacy_shape(cls, value: Any) -> Any:
        if isinstance(value, dict) and "analysis_targets" not in value and "primary_sku" in value:
            channels = value.get("channels") or ["JD"]
            return {
                **{key: item for key, item in value.items() if key not in {"primary_sku", "channels"}},
                "analysis_targets": [{"channel": channel, "primary_sku": value["primary_sku"]} for channel in channels],
            }
        return value

    @field_validator("analysis_targets")
    @classmethod
    def unique_targets(cls, value: list[AnalysisTarget]) -> list[AnalysisTarget]:
        result: list[AnalysisTarget] = []
        seen: set[tuple[str, str]] = set()
        for target in value:
            key = (target.channel, target.primary_sku)
            if key not in seen:
                result.append(target)
                seen.add(key)
        channels = [item.channel for item in result]
        if len(channels) != len(set(channels)):
            raise ValueError("Each channel may appear only once; provide one explicit SKU per channel")
        return result

    @property
    def channels(self) -> list[str]:
        return [item.channel for item in self.analysis_targets]

    @property
    def primary_sku(self) -> str:
        return self.analysis_targets[0].primary_sku

    def sku_for(self, channel: str) -> str:
        for target in self.analysis_targets:
            if target.channel == channel.upper():
                return target.primary_sku
        raise KeyError(channel)


class ConfirmedMapping(BaseModel):
    brand: str
    model: str

    @field_validator("brand")
    @classmethod
    def uppercase_brand(cls, value: str) -> str:
        return value.strip().upper()


class InsightClaim(BaseModel):
    model_config = ConfigDict(extra="forbid")

    claim_id: str = Field(pattern=r"^CLAIM-\d{3}$")
    claim_type: Literal["internal", "external", "integrated", "decision"]
    subject: Literal["PHILIPS", "BRAUN", "PANASONIC", "FLYCO", "OFFICIAL_EVENTS", "INTEGRATED", "DECISION"]
    text: str = Field(min_length=4, max_length=420)
    observation_ids: list[str] = Field(default_factory=list)
    external_evidence_ids: list[str] = Field(default_factory=list)
    evidence_level: Literal["OBSERVED", "SUPPORTED", "HYPOTHESIS", "NOT_ASSESSABLE"]
    alternative_explanations: list[str] = Field(default_factory=list)
    missing_data: list[str] = Field(default_factory=list)


class DecisionRecommendation(BaseModel):
    model_config = ConfigDict(extra="forbid")

    action: str = Field(min_length=8, max_length=900)
    channel: Literal["JD", "ALI", "OFFLINE"]
    timing: str = Field(min_length=2, max_length=300)
    proposed_range: str = Field(min_length=2, max_length=300)
    target_metrics: list[Literal["ASP", "Qty", "GMV", "Gross Margin"]] = Field(min_length=1)
    guardrails: list[str] = Field(min_length=1)
    stop_condition: str = Field(min_length=4, max_length=500)
    validation_design: str = Field(min_length=8, max_length=700)
    evidence_level: Literal["SUPPORTED", "HYPOTHESIS", "NOT_ASSESSABLE"]
    observation_ids: list[str] = Field(default_factory=list)
    external_evidence_ids: list[str] = Field(default_factory=list)
    missing_data: list[str] = Field(default_factory=list)


class InsightPackage(BaseModel):
    model_config = ConfigDict(extra="forbid")

    internal_drivers: list[InsightClaim] = Field(min_length=2, max_length=2)
    external_drivers: list[InsightClaim] = Field(min_length=4, max_length=4)
    integrated_conclusion: list[InsightClaim] = Field(min_length=3, max_length=3)
    decision_claim: InsightClaim
    decision: DecisionRecommendation

    @model_validator(mode="before")
    @classmethod
    def accept_ranked_claims(cls, value: Any) -> Any:
        if isinstance(value, dict) and "claims" in value and "internal_drivers" not in value:
            claims = [InsightClaim.model_validate(item) for item in value["claims"]]
            grouped = [
                *[item for item in claims if item.claim_type == "internal"],
                *[item for item in claims if item.claim_type == "external"],
                *[item for item in claims if item.claim_type == "integrated"],
                *[item for item in claims if item.claim_type == "decision"],
            ]
            if len(grouped) == 10:
                grouped = [item.model_copy(update={"claim_id": f"CLAIM-{index:03d}"}) for index, item in enumerate(grouped, 1)]
                return {"internal_drivers": grouped[:2], "external_drivers": grouped[2:6], "integrated_conclusion": grouped[6:9], "decision_claim": grouped[9], "decision": value.get("decision")}
        return value

    @model_validator(mode="after")
    def fixed_roles(self) -> "InsightPackage":
        self.internal_drivers = [item.model_copy(update={"claim_id": f"CLAIM-{index:03d}", "claim_type": "internal", "subject": "PHILIPS"}) for index, item in enumerate(self.internal_drivers, 1)]
        external_subjects = ["BRAUN", "PANASONIC", "FLYCO", "OFFICIAL_EVENTS"]
        self.external_drivers = [item.model_copy(update={"claim_id": f"CLAIM-{index:03d}", "claim_type": "external", "subject": subject}) for index, (item, subject) in enumerate(zip(self.external_drivers, external_subjects), 3)]
        self.integrated_conclusion = [item.model_copy(update={"claim_id": f"CLAIM-{index:03d}", "claim_type": "integrated", "subject": "INTEGRATED"}) for index, item in enumerate(self.integrated_conclusion, 7)]
        self.decision_claim = self.decision_claim.model_copy(update={"claim_id": "CLAIM-010", "claim_type": "decision", "subject": "DECISION"})
        return self

    @property
    def claims(self) -> list[InsightClaim]:
        return [*self.internal_drivers, *self.external_drivers, *self.integrated_conclusion, self.decision_claim]


class InternalInsightSection(BaseModel):
    model_config = ConfigDict(extra="forbid")
    claims: list[InsightClaim] = Field(min_length=2, max_length=2)


class ExternalInsightSection(BaseModel):
    model_config = ConfigDict(extra="forbid")
    claims: list[InsightClaim] = Field(min_length=4, max_length=4)


class IntegratedInsightSection(BaseModel):
    model_config = ConfigDict(extra="forbid")
    claims: list[InsightClaim] = Field(min_length=3, max_length=3)


class DecisionInsightSection(BaseModel):
    model_config = ConfigDict(extra="forbid")
    claim: InsightClaim
    decision: DecisionRecommendation


@dataclass(frozen=True)
class EvidenceRecord:
    evidence_id: str
    brand: str
    model: str
    period_label: str
    title: str
    url: str
    published_at: str | None
    retrieved_at: str
    source_domain: str
    official: bool
    snippet: str
    content_sha256: str


METRIC_SCHEMA = pa.DataFrameSchema(
    {
        "period": pa.Column(pa.DateTime), "period_label": pa.Column(str, Check.str_matches(r"^\d{2}-\d{2}$")),
        "brand": pa.Column(str, Check.str_length(min_value=1)), "model": pa.Column(str, Check.str_length(min_value=1)),
        "role": pa.Column(str, Check.isin(["PH", "Competitor"])), "price": pa.Column(float, Check.ge(0), nullable=True, coerce=True),
        "sales": pa.Column(float, Check.ge(0), nullable=True, coerce=True), "sales_value": pa.Column(float, Check.ge(0), nullable=True, coerce=True),
        "source_sheet": pa.Column(str, Check.str_length(min_value=1)), "mapping_state": pa.Column(str, Check.str_length(min_value=1)),
    }, strict=True, coerce=True,
)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_agnes_api_key(root: Path) -> str | None:
    """Load a local Agnes key without copying it into project artifacts or logs."""
    if os.getenv("AGNES_API_KEY", "").strip():
        return "environment"
    if os.getenv("AGNES_DISABLE_KEY_FILE", "").strip().lower() in {"1", "true", "yes"}:
        return None
    configured = os.getenv("AGNES_API_KEY_FILE", "").strip()
    candidates = [Path(configured)] if configured else [root / "api.txt", root.parent / "api.txt"]
    for candidate in candidates:
        if not candidate.is_file():
            continue
        raw = candidate.read_text(encoding="utf-8-sig").strip()
        token = raw
        if raw.startswith("{"):
            try:
                payload = json.loads(raw)
                token = str(payload.get("AGNES_API_KEY") or payload.get("api_key") or payload.get("key") or "").strip()
            except Exception:
                continue
        elif "=" in raw and "\n" not in raw:
            name, value = raw.split("=", 1)
            token = value.strip() if name.strip().upper() in {"AGNES_API_KEY", "API_KEY", "KEY"} else ""
        if not 20 <= len(token) <= 500 or any(character.isspace() for character in token):
            continue
        os.environ["AGNES_API_KEY"] = token
        return "local_key_file"
    return None


def rule_based_intent(prompt: str) -> AnalysisIntent:
    sku_matches = prompt_sku_matches(prompt)
    if not sku_matches:
        raise ValueError("未识别到 Philips SKU，例如 S3203/08")
    upper = prompt.upper()
    channel_hits: list[tuple[str, int]] = []
    for channel, aliases in CHANNEL_NAME_ALIASES.items():
        positions = [upper.find(alias.upper()) for alias in aliases if upper.find(alias.upper()) >= 0]
        if positions:
            channel_hits.append((channel, min(positions)))
    channels = [item[0] for item in sorted(channel_hits, key=lambda item: item[1])] or ["JD"]
    if len(channels) == 1 and len(sku_matches) > 1:
        expression = prompt[sku_matches[0].start():sku_matches[-1].end()]
        connector_residue = PHILIPS_SKU_PATTERN.sub("", expression)
        if re.sub(r"[\s/、,，;+&]+", "", connector_residue):
            raise ValueError("一次只能分析一个 Philips 产品；同一产品的组合型号必须使用斜杠连接。")
        product_group = " / ".join(dict.fromkeys(item.group(0).upper() for item in sku_matches))
        targets = [AnalysisTarget(channel=channels[0], primary_sku=product_group)]
    elif len(sku_matches) == 1:
        targets = [AnalysisTarget(channel=channel, primary_sku=sku_matches[0].group(0)) for channel in channels]
    else:
        if len(channels) != len(sku_matches):
            raise ValueError("多渠道请求必须为每个渠道明确提供一个 Philips SKU")
        remaining = list(sku_matches)
        targets = []
        for channel, position in sorted(channel_hits, key=lambda item: item[1]):
            nearest = min(remaining, key=lambda item: abs(item.start() - position))
            targets.append(AnalysisTarget(channel=channel, primary_sku=nearest.group(0)))
            remaining.remove(nearest)
    month_match = re.search(r"(\d{1,2})\s*个?月", prompt)
    return AnalysisIntent(analysis_targets=targets, months=int(month_match.group(1)) if month_match else 12)


class AgnesClient:
    def __init__(self, audit_path: Path | None = None, credential_source: str | None = None) -> None:
        self.api_key = os.getenv("AGNES_API_KEY", "").strip()
        self.base_url = os.getenv("AGNES_BASE_URL", "https://apihub.agnes-ai.com/v1").rstrip("/")
        self.model = os.getenv("AGNES_MODEL", "agnes-2.5-flash")
        self.timeout = float(os.getenv("AGNES_TIMEOUT_SECONDS", "90"))
        self.required = os.getenv("AGNES_REQUIRED", "false").strip().lower() in {"1", "true", "yes"}
        self.allow_fallback = os.getenv("AGNES_ALLOW_FALLBACK", "true").strip().lower() in {"1", "true", "yes"}
        self.response_format = os.getenv("AGNES_RESPONSE_FORMAT", "json_schema").strip().lower()
        self.audit_path = audit_path
        self.credential_source = credential_source or ("environment" if self.api_key else None)
        if self.audit_path is not None:
            self.audit_path.parent.mkdir(parents=True, exist_ok=True)
        self._lock = threading.Lock()
        self._status: dict[str, Any] = {
            "last_operation": None,
            "last_status": "not_called",
            "last_error": None,
            "last_success_at": None,
            "last_latency_ms": None,
            "request_count": 0,
            "fallback_count": 0,
        }

    @property
    def enabled(self) -> bool:
        return bool(self.api_key)

    def status(self) -> dict[str, Any]:
        with self._lock:
            return {
                "configured": self.enabled,
                "required": self.required,
                "allow_fallback": self.allow_fallback,
                "model": self.model if self.enabled else None,
                "credential_source": self.credential_source,
                "audit_enabled": self.audit_path is not None,
                **self._status,
            }

    def _update_status(self, **values: Any) -> None:
        with self._lock:
            self._status.update(values)

    def _mark_fallback(self, operation: str, error: Exception | None = None) -> None:
        with self._lock:
            self._status["last_operation"] = operation
            self._status["last_status"] = "fallback"
            self._status["last_error"] = str(error)[:500] if error else None
            self._status["fallback_count"] += 1
        self._append_audit(operation=operation, status="fallback", attempt=0, latency_ms=None, response_hash=None, schema_name=None, error=error)

    def _record_attempt(self, **values: Any) -> None:
        with self._lock:
            self._status["request_count"] += 1
            self._status.update(values)

    def _append_audit(self, *, operation: str, status: str, attempt: int, latency_ms: float | None, response_hash: str | None, schema_name: str | None, error: Exception | None) -> None:
        if self.audit_path is None:
            return
        event = {
            "timestamp": datetime.now(timezone.utc).isoformat(),
            "operation": operation,
            "status": status,
            "attempt": attempt,
            "model": self.model if self.enabled else None,
            "schema": schema_name,
            "response_sha256": response_hash,
            "latency_ms": latency_ms,
            "error": str(error)[:500] if error else None,
        }
        with self._lock:
            with self.audit_path.open("a", encoding="utf-8") as stream:
                stream.write(json.dumps(event, ensure_ascii=False) + "\n")

    def _chat_model(
        self,
        system: str,
        user: str,
        response_model: type[BaseModel],
        operation: str,
        semantic_validator: Callable[[BaseModel], BaseModel] | None = None,
        *,
        max_attempts: int = 3,
        request_timeout: float | None = None,
    ) -> BaseModel:
        if not self.enabled:
            raise RuntimeError("AGNES_API_KEY is not configured")
        headers = {"Authorization": f"Bearer {self.api_key}", "Content-Type": "application/json"}
        last_error: Exception | None = None
        feedback = ""
        for attempt in range(max_attempts):
            messages = [{"role": "system", "content": system}, {"role": "user", "content": user + feedback}]
            payload: dict[str, Any] = {"model": self.model, "messages": messages, "temperature": 0, "max_tokens": 4096, "stream": False}
            if self.response_format == "json_schema":
                payload["response_format"] = {
                    "type": "json_schema",
                    "json_schema": {"name": response_model.__name__, "strict": True, "schema": response_model.model_json_schema()},
                }
            elif self.response_format == "json_object":
                payload["response_format"] = {"type": "json_object"}
            started = time.perf_counter()
            content = ""
            try:
                response = httpx.post(
                    f"{self.base_url}/chat/completions",
                    headers=headers,
                    json=payload,
                    timeout=request_timeout if request_timeout is not None else self.timeout,
                )
                response.raise_for_status()
                content = str(response.json()["choices"][0]["message"]["content"]).strip()
                if content.startswith("```"):
                    content = re.sub(r"^```(?:json)?\s*|\s*```$", "", content, flags=re.I | re.S)
                parsed = json.loads(content)
                if not isinstance(parsed, dict):
                    raise ValueError("Agnes response must be a JSON object")
                validated = response_model.model_validate(parsed)
                if semantic_validator is not None:
                    validated = semantic_validator(validated)
                latency = round((time.perf_counter() - started) * 1000, 1)
                response_hash = hashlib.sha256(content.encode("utf-8")).hexdigest()
                self._record_attempt(
                    last_operation=operation,
                    last_status="success",
                    last_error=None,
                    last_success_at=datetime.now(timezone.utc).isoformat(),
                    last_latency_ms=latency,
                )
                self._append_audit(operation=operation, status="success", attempt=attempt + 1, latency_ms=latency, response_hash=response_hash, schema_name=response_model.__name__, error=None)
                return validated
            except Exception as exc:
                last_error = exc
                latency = round((time.perf_counter() - started) * 1000, 1)
                response_hash = hashlib.sha256(content.encode("utf-8")).hexdigest() if content else None
                self._record_attempt(
                    last_operation=operation,
                    last_status="retrying" if attempt < max_attempts - 1 else "failed",
                    last_error=str(exc)[:500],
                    last_latency_ms=latency,
                )
                self._append_audit(operation=operation, status="retrying" if attempt < max_attempts - 1 else "failed", attempt=attempt + 1, latency_ms=latency, response_hash=response_hash, schema_name=response_model.__name__, error=exc)
                feedback = f"\n\nYour previous output failed schema validation: {str(exc)[:800]}. Return a corrected JSON object only."
        raise RuntimeError(f"Agnes request failed after retries: {last_error}")

    def parse_intent(self, prompt: str, *, safe_local_fallback: bool = False) -> AnalysisIntent:
        exact_prompt_intent = self._require_single_channel(rule_based_intent(prompt))
        if not self.enabled:
            if self.required:
                raise RuntimeError("AGNES_API_KEY is required but not configured")
            self._mark_fallback("intent")
            return exact_prompt_intent
        system = "Return one JSON object matching AnalysisIntent exactly. analysis_targets must contain exactly one target: the employee chooses JD, ALI, or OFFLINE, never multiple channels in one run. Preserve every member of a same-channel compound Philips product group such as YQ660/02/PQ663/02 in the same primary_sku string. Never infer cross-channel aliases. Preserve slashes. Defaults: 12 months, ASP and Unit, one_model_per_brand, key_sku_list_first, sku_insight_v20. Never add fields."
        try:
            recognized = AnalysisIntent.model_validate(
                self._chat_model(
                    system,
                    prompt,
                    AnalysisIntent,
                    "intent",
                    max_attempts=1 if safe_local_fallback else 3,
                    request_timeout=min(self.timeout, 20.0) if safe_local_fallback else None,
                )
            )
            return recognized.model_copy(update={"analysis_targets": exact_prompt_intent.analysis_targets})
        except Exception as exc:
            if self.required or (not self.allow_fallback and not safe_local_fallback):
                raise
            self._mark_fallback("intent", exc)
            return exact_prompt_intent

    @staticmethod
    def _require_single_channel(intent: AnalysisIntent) -> AnalysisIntent:
        if len(intent.analysis_targets) != 1:
            raise ValueError("一次只能分析一个销售渠道，请从 JD、ALI 或 Offline 中选择一个。")
        return intent

    def recognize_semantic_plan(self, manifest: WorkbookStructureManifest, draft: WorkbookSemanticPlan | None) -> WorkbookSemanticPlan:
        if not self.enabled:
            if draft is None:
                raise RuntimeError("Agnes is required to recognize an unfamiliar workbook layout")
            self._mark_fallback("semantic_plan")
            return draft
        system = "Classify workbook labels into a WorkbookSemanticPlan. Use only cells, raw_text, sheet names, and ranges supplied by the user. Never invent a cell, brand, channel, SKU, number, price, unit, SQL, or mapping. Channel groups must be JD, ALI, and OFFLINE, remain non-overlapping, and cite the exact channel/header cells. Mark unfamiliar or low-confidence labels as novel and set review_required. Return JSON only."
        payload = {"manifest": manifest.model_dump(), "deterministic_draft": draft.model_dump() if draft else None}
        plan = WorkbookSemanticPlan.model_validate(
            self._chat_model(system, json.dumps(payload, ensure_ascii=False), WorkbookSemanticPlan, "semantic_plan")
        )
        return plan.model_copy(update={"source": "agnes", "recognizer_model": self.model})

    @staticmethod
    def _conservatize_insight_package(package: InsightPackage, facts: dict[str, Any]) -> InsightPackage:
        missing_defaults = [str(item) for item in facts.get("missing_driver_inputs", [])] or ["verified driver evidence"]
        allowed_evidence = {item["evidence_id"] for item in facts.get("external_evidence", [])}
        causal_markers = ("导致", "由于", "归因于", "驱动", "证明", "反映", "表明", "说明", "模式成立", "caused", "because", "resulted in", "proves")

        def conservative_claim(claim: InsightClaim) -> InsightClaim:
            update: dict[str, Any] = {}
            valid_external = [item for item in claim.external_evidence_ids if item in allowed_evidence]
            already_conservative = any(
                marker in claim.text
                for marker in (
                    "不能证明外部原因或因果",
                    "不能把价格或销量变化归因于外部活动",
                    "不得把销量变化归因于",
                )
            )
            if claim.evidence_level == "SUPPORTED" and not valid_external:
                update.update(
                    evidence_level="HYPOTHESIS",
                    text="未检索到与分析月份对齐的白名单官方证据，当前不能把价格或销量变化归因于外部活动。",
                    external_evidence_ids=[],
                    alternative_explanations=claim.alternative_explanations or ["promotion", "traffic", "inventory"],
                    missing_data=claim.missing_data or ["verified official event evidence"],
                )
            elif not valid_external and not already_conservative and any(marker in claim.text.lower() for marker in causal_markers):
                lowered = claim.text.lower()
                cut = min(lowered.find(marker) for marker in causal_markers if marker in lowered)
                descriptive_prefix = claim.text[:cut].rstrip(" ，,；;。.")
                update.update(
                    evidence_level="OBSERVED" if claim.observation_ids else "HYPOTHESIS",
                    text=f"{descriptive_prefix}；该共同变化仅为观察，缺少官方事件及促销、流量、库存证据，不能证明外部原因或因果。",
                    external_evidence_ids=[],
                    alternative_explanations=claim.alternative_explanations or ["promotion", "traffic", "inventory"],
                    missing_data=claim.missing_data or missing_defaults,
                )
            elif claim.evidence_level == "HYPOTHESIS" and not (claim.alternative_explanations or claim.missing_data):
                update["missing_data"] = missing_defaults
            elif claim.evidence_level == "NOT_ASSESSABLE" and not claim.missing_data:
                update["missing_data"] = missing_defaults
            return claim.model_copy(update=update) if update else claim

        internal = [conservative_claim(item) for item in package.internal_drivers]
        external = [conservative_claim(item) for item in package.external_drivers]
        integrated = [conservative_claim(item) for item in package.integrated_conclusion]
        decision_claim = conservative_claim(package.decision_claim)
        decision = package.decision
        decision_update: dict[str, Any] = {}
        valid_decision_evidence = [item for item in decision.external_evidence_ids if item in allowed_evidence]
        if decision.evidence_level == "SUPPORTED" and not valid_decision_evidence:
            decision_update.update(evidence_level="HYPOTHESIS", external_evidence_ids=[], missing_data=decision.missing_data or missing_defaults)
        elif decision.evidence_level in {"HYPOTHESIS", "NOT_ASSESSABLE"} and not decision.missing_data:
            decision_update["missing_data"] = missing_defaults
        essential_missing = {item.lower().replace("_", " ") for item in [*decision.missing_data, *missing_defaults]}
        if {"promotion", "traffic", "inventory", "gross margin"} & essential_missing:
            primary_sku = str(facts.get("primary_sku", "the selected SKU"))
            decision_update.update(
                action=f"先核查 {primary_sku} 的促销、流量、库存与毛利，再在当前渠道进行可回滚的小范围价格验证。",
                timing="完成缺失数据核查并确认当前竞品映射仍有效后",
                proposed_range="毛利和促销约束缺失，当前不设定具体价格区间",
                target_metrics=["ASP", "Qty", "GMV", "Gross Margin"],
                guardrails=["销量、GMV 与毛利必须同时评估", "促销、流量或库存口径不可比时停止", "竞品映射变化时重新生成分析"],
                stop_condition="GMV 或毛利未达到业务门槛，或发现促销、流量、库存不可比时停止并回滚。",
                validation_design="在同一渠道选择可比月份开展受控、可回滚的价格测试，并同步记录 ASP、Qty、GMV、毛利、促销、流量与库存。",
                evidence_level="HYPOTHESIS",
                external_evidence_ids=[],
                missing_data=missing_defaults,
            )
            decision_claim = decision_claim.model_copy(update={
                "text": f"下一步先补齐 {primary_sku} 的促销、流量、库存和毛利证据，再进行不预设具体价格区间的受控测试；以 Qty、GMV 与毛利作为共同决策门槛。",
                "evidence_level": "HYPOTHESIS",
                "external_evidence_ids": [],
                "alternative_explanations": decision_claim.alternative_explanations or ["promotion", "traffic", "inventory"],
                "missing_data": missing_defaults,
            })
        if decision_update:
            decision = decision.model_copy(update=decision_update)
        return package.model_copy(update={"internal_drivers": internal, "external_drivers": external, "integrated_conclusion": integrated, "decision_claim": decision_claim, "decision": decision})

    @staticmethod
    def _validate_insight_package(package: InsightPackage, facts: dict[str, Any]) -> InsightPackage:
        allowed_observations = {item["observation_id"] for item in facts.get("observations", [])}
        allowed_evidence = {item["evidence_id"] for item in facts.get("external_evidence", [])}
        observations_by_id = {item["observation_id"]: item for item in facts.get("observations", [])}
        expected_claim_ids = [f"CLAIM-{index:03d}" for index in range(1, 11)]
        if [item.claim_id for item in package.claims] != expected_claim_ids:
            raise ValueError("Claims must be ranked CLAIM-001 through CLAIM-010")
        claim_types = [item.claim_type for item in package.claims]
        if claim_types.count("internal") < 2 or claim_types.count("external") < 2 or claim_types.count("integrated") < 2 or claim_types[-1] != "decision":
            raise ValueError("Insight package does not preserve internal → external → integrated → decision coverage")
        for claim in package.claims:
            AgnesClient._require_chinese_narrative(claim.text, f"{claim.claim_id}.text")
            if set(claim.observation_ids) - allowed_observations or set(claim.external_evidence_ids) - allowed_evidence:
                raise ValueError(f"Unknown evidence identifier in {claim.claim_id}")
            cited_observations = [observations_by_id[item] for item in claim.observation_ids]
            if claim.subject in {"BRAUN", "PANASONIC", "FLYCO"}:
                cited_brands = {str(item.get("brand", "")).upper() for item in cited_observations if item.get("brand")}
                if any(brand not in {claim.subject, "PHILIPS"} for brand in cited_brands):
                    raise ValueError(f"{claim.claim_id} cites observations outside subject {claim.subject}")
                if claim.observation_ids and claim.subject not in cited_brands:
                    raise ValueError(f"{claim.claim_id} cites observations but none belong to subject {claim.subject}")
            if claim.subject == "PHILIPS" and any(str(item.get("brand", "")).upper() != "PHILIPS" for item in cited_observations if item.get("brand")):
                raise ValueError(f"{claim.claim_id} internal evidence is not Philips-only")
            AgnesClient._validate_claim_numbers(claim, cited_observations, facts)
            if claim.evidence_level == "OBSERVED" and not claim.observation_ids:
                raise ValueError(f"OBSERVED claim {claim.claim_id} requires observation IDs")
            if claim.evidence_level == "SUPPORTED" and not claim.external_evidence_ids:
                raise ValueError(f"SUPPORTED claim {claim.claim_id} requires official evidence IDs")
            if claim.evidence_level == "NOT_ASSESSABLE" and not claim.missing_data:
                raise ValueError(f"NOT_ASSESSABLE claim {claim.claim_id} must name missing data")
            if claim.evidence_level == "HYPOTHESIS" and not (claim.alternative_explanations or claim.missing_data):
                raise ValueError(f"HYPOTHESIS claim {claim.claim_id} must disclose uncertainty")
        decision = package.decision
        decision_fields = {
            "action": decision.action,
            "timing": decision.timing,
            "proposed_range": decision.proposed_range,
            "stop_condition": decision.stop_condition,
            "validation_design": decision.validation_design,
        }
        for field_name, value in decision_fields.items():
            AgnesClient._require_chinese_narrative(value, f"decision.{field_name}")
        for index, value in enumerate(decision.guardrails):
            AgnesClient._require_chinese_narrative(value, f"decision.guardrails[{index}]")
        if decision.channel != facts.get("channel"):
            raise ValueError("Decision channel does not match the analyzed channel")
        if set(decision.observation_ids) - allowed_observations or set(decision.external_evidence_ids) - allowed_evidence:
            raise ValueError("Decision cites an unknown evidence identifier")
        decision_observations = [observations_by_id[item] for item in decision.observation_ids]
        decision_text = "；".join([
            decision.action,
            decision.timing,
            decision.proposed_range,
            *decision.guardrails,
            decision.stop_condition,
            decision.validation_design,
        ])
        AgnesClient._validate_claim_numbers(
            package.decision_claim.model_copy(
                update={
                    "text": decision_text,
                    "observation_ids": decision.observation_ids,
                    "external_evidence_ids": decision.external_evidence_ids,
                }
            ),
            decision_observations,
            facts,
        )
        if decision.evidence_level == "SUPPORTED" and not decision.external_evidence_ids:
            raise ValueError("SUPPORTED decision requires official evidence")
        if decision.evidence_level in {"HYPOTHESIS", "NOT_ASSESSABLE"} and not decision.missing_data:
            raise ValueError("Unproven decision must disclose missing data")
        return package

    @staticmethod
    def _require_chinese_narrative(value: str, field_name: str) -> None:
        cjk_count = len(re.findall(r"[\u3400-\u9fff]", value))
        latin_count = len(re.findall(r"[A-Za-z]", value))
        if cjk_count < 4 or latin_count > max(24, cjk_count * 3):
            raise ValueError(f"{field_name} must be written in Chinese; Latin text is limited to brand, SKU and metric names")

    @staticmethod
    def _validate_claim_numbers(claim: InsightClaim, cited_observations: list[dict[str, Any]], facts: dict[str, Any]) -> None:
        allowed_numbers: list[float] = []

        def collect(value: Any) -> None:
            if isinstance(value, bool) or value is None:
                return
            if isinstance(value, (int, float)) and math.isfinite(float(value)):
                allowed_numbers.append(float(value))
            elif isinstance(value, dict):
                for item in value.values():
                    collect(item)
            elif isinstance(value, list):
                for item in value:
                    collect(item)

        collect(cited_observations)
        scrubbed = claim.text
        labels = [
            str(facts.get("primary_sku", "")),
            *[str(item) for item in facts.get("period_scope", [])],
            *claim.observation_ids,
            *claim.external_evidence_ids,
        ]
        for item in cited_observations:
            labels.extend([str(item.get("brand", "")), str(item.get("model", ""))])
        for label in sorted({item for item in labels if item}, key=len, reverse=True):
            scrubbed = re.sub(re.escape(label), "", scrubbed, flags=re.I)
        scrubbed = re.sub(r"\b(?:level\s*[123]|top\s*10)\b", "", scrubbed, flags=re.I)
        tokens = re.findall(r"(?<![A-Za-z])[-+−]?\d[\d,]*(?:\.\d+)?", scrubbed)
        unsupported: list[str] = []
        for token in tokens:
            normalized = token.replace(",", "").replace("−", "-")
            try:
                number = float(normalized)
            except ValueError:
                continue
            explicitly_signed = normalized.startswith(("+", "-"))
            grounded = any(
                abs(number - allowed) <= max(0.11, abs(allowed) * 0.0006)
                or (not explicitly_signed and abs(abs(number) - abs(allowed)) <= max(0.11, abs(allowed) * 0.0006))
                for allowed in allowed_numbers
            )
            if not grounded:
                unsupported.append(token)
        if unsupported:
            raise ValueError(f"{claim.claim_id} contains numbers not grounded in its cited observations: {unsupported[:6]}")

    def synthesize_insights(self, facts: dict[str, Any], fallback: InsightPackage) -> InsightPackage:
        if not self.enabled:
            if self.required:
                raise RuntimeError("AGNES_API_KEY is required but not configured")
            self._mark_fallback("insights")
            return fallback
        common = """You are one stage in a governed SKU Insight Agent. Return exactly one JSON object matching the requested schema and no prose. Every user-visible narrative field must be written in concise Simplified Chinese; Latin characters are allowed only for brand names, product SKUs, metric names such as ASP/Qty/GMV, and evidence IDs. Each claim text must stay within 60 Chinese characters where possible and state the conclusion before its limitation. Use only supplied observation_id and evidence_id values. Do not put digits, percentages, quantities, dates, ranges or calculated values in any claim text; exact numbers are rendered separately by the deterministic metrics layer. Describe only direction, relationship, risk, opportunity, uncertainty and a testable interpretation. A product SKU may be named exactly as supplied. OBSERVED requires observation IDs. SUPPORTED requires verified official external evidence IDs. HYPOTHESIS and NOT_ASSESSABLE must disclose alternatives or missing data. Never invent an event, URL, SKU, number, causal link, threshold or timing. Correlation and coincident movement are not causation. The local policy layer assigns final claim IDs, roles, subjects and evidence bindings, so focus on grounded qualitative content."""
        base = {
            "channel": facts.get("channel"),
            "primary_sku": facts.get("primary_sku"),
            "period_scope": facts.get("period_scope", []),
            "quality_warnings": facts.get("quality_warnings", []),
            "missing_driver_inputs": facts.get("missing_driver_inputs", []),
            "method_constraints": facts.get("method_constraints", []),
        }
        governed_observations = [
            item for item in facts.get("observations", [])
            if item.get("role") == "PH" or str(item.get("brand", "")).upper() in {"BRAUN", "PANASONIC", "FLYCO"}
        ]

        def compose(
            *,
            internal: list[InsightClaim] | None = None,
            external: list[InsightClaim] | None = None,
            integrated: list[InsightClaim] | None = None,
            decision_section: DecisionInsightSection | None = None,
        ) -> InsightPackage:
            all_observations = governed_observations
            primary_ids = [item["observation_id"] for item in all_observations if item.get("role") == "PH"]
            all_observation_ids = [item["observation_id"] for item in all_observations]
            allowed_evidence_ids = {item["evidence_id"] for item in facts.get("external_evidence", [])}

            normalized_internal = [item.model_copy(update={"observation_ids": primary_ids, "external_evidence_ids": []}) for item in (internal or fallback.internal_drivers)]
            normalized_external: list[InsightClaim] = []
            for item, subject in zip(external or fallback.external_drivers, ("BRAUN", "PANASONIC", "FLYCO", "OFFICIAL_EVENTS")):
                subject_ids: list[str] = []
                if subject == "OFFICIAL_EVENTS":
                    cited = primary_ids
                else:
                    subject_ids = [
                        observation["observation_id"]
                        for observation in all_observations
                        if str(observation.get("brand", "")).upper() == subject
                    ]
                    cited = [*primary_ids, *subject_ids] if subject_ids else []
                update: dict[str, Any] = {
                    "observation_ids": list(dict.fromkeys(cited)),
                    "external_evidence_ids": [evidence_id for evidence_id in item.external_evidence_ids if evidence_id in allowed_evidence_ids],
                }
                if subject != "OFFICIAL_EVENTS" and not subject_ids:
                    update.update({
                        "text": f"{subject} 在当前渠道没有已验证的映射型号或有效月度数据，因此无法评估其价格与销量变化。",
                        "evidence_level": "NOT_ASSESSABLE",
                        "observation_ids": [],
                        "external_evidence_ids": [],
                        "alternative_explanations": [],
                        "missing_data": [f"{subject} same-channel mapping and metrics"],
                    })
                elif subject == "OFFICIAL_EVENTS" and not allowed_evidence_ids:
                    update.update({
                        "text": "未检索到与分析月份对齐的白名单官方证据，当前不能把价格或销量变化归因于外部活动。",
                        "evidence_level": "HYPOTHESIS",
                        "external_evidence_ids": [],
                        "alternative_explanations": ["promotion", "traffic", "inventory"],
                        "missing_data": ["verified official event evidence"],
                    })
                normalized_external.append(item.model_copy(update=update))
            normalized_integrated = [item.model_copy(update={
                "observation_ids": all_observation_ids,
                "external_evidence_ids": [evidence_id for evidence_id in item.external_evidence_ids if evidence_id in allowed_evidence_ids],
            }) for item in (integrated or fallback.integrated_conclusion)]
            selected_decision = decision_section or DecisionInsightSection(claim=fallback.decision_claim, decision=fallback.decision)
            normalized_decision = DecisionInsightSection(
                claim=selected_decision.claim.model_copy(update={
                    "observation_ids": all_observation_ids,
                    "external_evidence_ids": [evidence_id for evidence_id in selected_decision.claim.external_evidence_ids if evidence_id in allowed_evidence_ids],
                }),
                decision=selected_decision.decision.model_copy(update={
                    "observation_ids": all_observation_ids,
                    "external_evidence_ids": [evidence_id for evidence_id in selected_decision.decision.external_evidence_ids if evidence_id in allowed_evidence_ids],
                }),
            )
            package = InsightPackage(
                internal_drivers=normalized_internal,
                external_drivers=normalized_external,
                integrated_conclusion=normalized_integrated,
                decision_claim=normalized_decision.claim,
                decision=normalized_decision.decision,
            )
            return self._validate_insight_package(self._conservatize_insight_package(package, facts), facts)

        try:
            internal_payload = {
                **base,
                "observations": [item for item in facts.get("observations", []) if item.get("role") == "PH"],
            }
            internal = InternalInsightSection.model_validate(self._chat_model(
                common + " Produce exactly two Philips internal claims. Claim 1 covers Level 1 ASP, Qty and GMV movement. Claim 2 covers Level 2 price/Qty quadrant, volatility, anomaly or elasticity limits. Do not discuss competitors or external causes.",
                json.dumps(internal_payload, ensure_ascii=False, default=str),
                InternalInsightSection,
                "insights_internal",
                lambda section: InternalInsightSection(claims=compose(internal=InternalInsightSection.model_validate(section).claims).internal_drivers),
            ))

            external_payload = {
                **base,
                "observations": governed_observations,
                "external_evidence": facts.get("external_evidence", []),
            }
            external = ExternalInsightSection.model_validate(self._chat_model(
                common + " Produce exactly four Level 3 external claims in this order: BRAUN, PANASONIC, FLYCO, OFFICIAL_EVENTS. For each competitor use only that brand's observations plus Philips observations. If a brand or verified official event is absent, return NOT_ASSESSABLE and name the missing data. Do not include YOOSE.",
                json.dumps(external_payload, ensure_ascii=False, default=str),
                ExternalInsightSection,
                "insights_external",
                lambda section: ExternalInsightSection(claims=compose(external=ExternalInsightSection.model_validate(section).claims).external_drivers),
            ))

            integrated_payload = {
                **base,
                "observations": governed_observations,
                "external_evidence": facts.get("external_evidence", []),
                "validated_internal_claims": [item.model_dump() for item in internal.claims],
                "validated_external_claims": [item.model_dump() for item in external.claims],
            }
            integrated = IntegratedInsightSection.model_validate(self._chat_model(
                common + " Produce exactly three integrated Monthly Report claims: market change, key risk, and key opportunity. Synthesize the supplied validated internal and external claims without adding new facts or causal certainty. State limitations where promotion, traffic, inventory or margin inputs are missing.",
                json.dumps(integrated_payload, ensure_ascii=False, default=str),
                IntegratedInsightSection,
                "insights_integrated",
                lambda section: IntegratedInsightSection(claims=compose(internal=internal.claims, external=external.claims, integrated=IntegratedInsightSection.model_validate(section).claims).integrated_conclusion),
            ))

            decision_payload = {
                **base,
                "validated_internal_claims": [item.model_dump() for item in internal.claims],
                "validated_external_claims": [item.model_dump() for item in external.claims],
                "validated_integrated_claims": [item.model_dump() for item in integrated.claims],
            }
            decision_section = DecisionInsightSection.model_validate(self._chat_model(
                common + " Produce one final decision claim and one structured recommendation. Preserve the exact supplied channel. Cover action, timing, proposed range, ASP/Qty/GMV/Gross Margin targets, guardrails, stop condition and validation design. If margin, promotion, traffic or inventory evidence is missing, recommend a reversible controlled test and do not set numeric price ranges or thresholds.",
                json.dumps(decision_payload, ensure_ascii=False, default=str),
                DecisionInsightSection,
                "insights_decision",
                lambda section: DecisionInsightSection(
                    claim=compose(
                        internal=internal.claims,
                        external=external.claims,
                        integrated=integrated.claims,
                        decision_section=DecisionInsightSection.model_validate(section),
                    ).decision_claim,
                    decision=compose(
                        internal=internal.claims,
                        external=external.claims,
                        integrated=integrated.claims,
                        decision_section=DecisionInsightSection.model_validate(section),
                    ).decision,
                ),
            ))
            result = compose(
                internal=internal.claims,
                external=external.claims,
                integrated=integrated.claims,
                decision_section=decision_section,
            )
            self._update_status(last_operation="insights", last_status="success", last_error=None)
            return result
        except Exception as exc:
            if self.required or not self.allow_fallback:
                raise
            self._mark_fallback("insights", exc)
            return fallback

    def synthesize_claims(self, facts: dict[str, Any], fallback: list[InsightClaim]) -> list[InsightClaim]:
        return self.synthesize_insights(facts, InsightPackage(claims=fallback, decision=deterministic_decision(facts))).claims


def pearson(xs: Iterable[float], ys: Iterable[float]) -> float | None:
    pairs = [(float(x), float(y)) for x, y in zip(xs, ys) if x is not None and y is not None]
    if len(pairs) < 3:
        return None
    xv, yv = zip(*pairs); mx, my = statistics.mean(xv), statistics.mean(yv)
    top = sum((x - mx) * (y - my) for x, y in pairs)
    bottom = math.sqrt(sum((x - mx) ** 2 for x in xv) * sum((y - my) ** 2 for y in yv))
    return top / bottom if bottom else None


def rank(values: list[float]) -> list[float]:
    indexed = sorted(enumerate(values), key=lambda item: item[1]); result = [0.0] * len(values); index = 0
    while index < len(indexed):
        end = index
        while end + 1 < len(indexed) and indexed[end + 1][1] == indexed[index][1]: end += 1
        average = (index + end + 2) / 2
        for position in range(index, end + 1): result[indexed[position][0]] = average
        index = end + 1
    return result


def validate_metrics(records: list[MonthlyMetric], sku: str, months: int, quality_config: dict[str, Any]) -> dict[str, Any]:
    rows = []
    for record in records:
        value = record.to_dict(); value.pop("source_cells"); rows.append(value)
    frame = pd.DataFrame(rows); frame["period"] = pd.to_datetime(frame["period"])
    validated = METRIC_SCHEMA.validate(frame, lazy=True)
    if validated.duplicated(["period", "brand", "model"]).any():
        raise WorkbookContractError("DUPLICATE_CANONICAL_KEY", "Duplicate period/brand/model records exist")
    primary = validated[validated["role"] == "PH"]
    if len(primary) != months or primary["price"].isna().any() or primary["sales"].isna().any():
        raise WorkbookContractError("PRIMARY_METRIC_INCOMPLETE", f"Primary SKU {sku} does not have complete ASP and Unit for {months} months")
    if len(set(primary["period_label"])) != months:
        raise WorkbookContractError("PRIMARY_MONTH_INCOMPLETE", "Primary SKU month set is incomplete")
    scales = [float(item) for item in quality_config["allowed_value_scales"]]; tolerance = float(quality_config["value_reconciliation_relative_tolerance"]); reconciliation = []
    for record in records:
        if record.price is None or record.sales is None or record.sales_value is None: continue
        expected = record.price * record.sales
        if expected == 0: continue
        scale = min(scales, key=lambda item: abs(record.sales_value - expected * item) / max(abs(record.sales_value), abs(expected * item), 1))
        error = abs(record.sales_value - expected * scale) / max(abs(record.sales_value), abs(expected * scale), 1)
        row = {"period": record.period_label, "brand": record.brand, "model": record.model, "scale": scale, "relative_error": error}; reconciliation.append(row)
        if error > tolerance: raise WorkbookContractError("VALUE_RECONCILIATION_FAILED", f"Value does not reconcile with ASP × Unit for {record.brand} {record.model} {record.period_label}", row)
    warnings = []; threshold = float(quality_config["competitor_minimum_coverage"])
    for brand, group in validated[validated["role"] == "Competitor"].groupby("brand"):
        coverage = float((group["price"].notna() & group["sales"].notna()).mean())
        if coverage < threshold: warnings.append({"code": "LOW_COMPETITOR_COVERAGE", "brand": brand, "coverage": coverage, "threshold": threshold})
    return {"status": "passed", "row_count": len(validated), "primary_months": months, "warnings": warnings, "reconciliation": reconciliation}


def percent_change(current: float | None, previous: float | None) -> float | None:
    return (float(current) / float(previous) - 1) * 100 if current is not None and previous not in (None, 0) else None


def elasticity_pattern(asp_change: float | None, qty_change: float | None) -> str:
    if asp_change is None or qty_change is None: return "Not assessable"
    asp = "Price ↑" if asp_change > 0.1 else "Price ↓" if asp_change < -0.1 else "Price stable"
    qty = "Qty ↑" if qty_change > 0.1 else "Qty ↓" if qty_change < -0.1 else "Qty stable"
    return f"{asp} / {qty}"


def summarize_series(records: list[MonthlyMetric]) -> dict[str, Any]:
    valid = sorted((item for item in records if item.price is not None and item.sales is not None), key=lambda item: item.period)
    if len(valid) < 2: return {"valid_months": len(valid), "status": "insufficient"}
    prices = [float(item.price) for item in valid]; sales = [float(item.sales) for item in valid]; gmvs = [float(item.sales_value) if item.sales_value is not None else float(item.price) * float(item.sales) for item in valid]; peak = max(valid, key=lambda item: item.sales or 0); latest, previous = valid[-1], valid[-2]
    prior_year = next((item for item in valid if item.period.year == latest.period.year - 1 and item.period.month == latest.period.month), None)
    asp_lm = percent_change(latest.price, previous.price); qty_lm = percent_change(latest.sales, previous.sales); gmv_lm = percent_change(latest.sales_value, previous.sales_value)
    qty_changes = [percent_change(valid[index].sales, valid[index - 1].sales) for index in range(1, len(valid))]; finite_changes = [abs(item) for item in qty_changes if item is not None]
    anomaly_threshold = max(25.0, statistics.median(finite_changes) * 3 if finite_changes else 25.0)
    material_moves = []
    for index in range(1, len(valid)):
        price_move = percent_change(valid[index].price, valid[index - 1].price)
        next_qty_move = percent_change(valid[index + 1].sales, valid[index].sales) if index + 1 < len(valid) else None
        if price_move is not None and abs(price_move) >= 5:
            material_moves.append({"period": valid[index].period_label, "asp_change_pct": price_move, "next_month_qty_change_pct": next_qty_move})
    min_price = min(valid, key=lambda item: item.price if item.price is not None else math.inf); max_price = max(valid, key=lambda item: item.price if item.price is not None else -math.inf)
    min_sales = min(valid, key=lambda item: item.sales if item.sales is not None else math.inf); max_sales = max(valid, key=lambda item: item.sales if item.sales is not None else -math.inf)
    return {"valid_months": len(valid), "first_period": valid[0].period_label, "last_period": latest.period_label, "price_change_pct": percent_change(prices[-1], prices[0]), "sales_change_pct": percent_change(sales[-1], sales[0]), "gmv_change_pct": percent_change(gmvs[-1], gmvs[0]), "latest_asp": latest.price, "latest_qty": latest.sales, "latest_gmv": latest.sales_value, "asp_lm_change_pct": asp_lm, "qty_lm_change_pct": qty_lm, "gmv_lm_change_pct": gmv_lm, "asp_ly_change_pct": percent_change(latest.price, prior_year.price) if prior_year else None, "qty_ly_change_pct": percent_change(latest.sales, prior_year.sales) if prior_year else None, "gmv_ly_change_pct": percent_change(latest.sales_value, prior_year.sales_value) if prior_year else None, "elasticity_pattern": elasticity_pattern(asp_lm, qty_lm), "anomaly": abs(qty_lm or 0) > anomaly_threshold, "anomaly_threshold_pct": anomaly_threshold, "pearson": pearson(prices, sales), "spearman": pearson(rank(prices), rank(sales)), "price_leads_sales_1m": pearson(prices[:-1], sales[1:]) if len(valid) >= 4 else None, "asp_volatility_cv_pct": statistics.pstdev(prices) / statistics.mean(prices) * 100 if statistics.mean(prices) else None, "qty_volatility_cv_pct": statistics.pstdev(sales) / statistics.mean(sales) * 100 if statistics.mean(sales) else None, "median_price": statistics.median(prices), "min_asp": min_price.price, "min_asp_period": min_price.period_label, "max_asp": max_price.price, "max_asp_period": max_price.period_label, "min_qty": min_sales.sales, "min_qty_period": min_sales.period_label, "max_qty": max_sales.sales, "max_qty_period": max_sales.period_label, "peak_sales_period": peak.period_label, "peak_sales": peak.sales, "material_price_moves": material_moves}


def build_observations(records: list[MonthlyMetric], mappings: list[MappingItem]) -> list[dict[str, Any]]:
    observations = []; index = 1
    for mapping in mappings:
        summary = summarize_series([item for item in records if normalize(item.brand) == normalize(mapping.brand) and normalize(item.model) == normalize(mapping.model)])
        observations.append({"observation_id": f"OBS-{index:03d}", "brand": mapping.brand, "model": mapping.model, "role": mapping.role, **summary}); index += 1
    primary = next((item for item in observations if item["role"] == "PH"), None)
    if primary and primary.get("median_price"):
        for competitor in [item for item in observations if item["role"] == "Competitor" and item.get("median_price")]:
            observations.append({"observation_id": f"OBS-{index:03d}", "type": "relative_price", "brand": competitor["brand"], "model": competitor["model"], "ph_price_premium_pct": (primary["median_price"] / competitor["median_price"] - 1) * 100}); index += 1
    return observations


def build_pricing_analysis_framework(observations: list[dict[str, Any]], package: InsightPackage, evidence: list[EvidenceRecord]) -> dict[str, Any]:
    primary = next((item for item in observations if item.get("role") == "PH"), {})
    competitors = {str(item.get("brand", "")).upper(): item for item in observations if item.get("role") == "Competitor"}
    relative = {str(item.get("brand", "")).upper(): item for item in observations if item.get("type") == "relative_price"}
    pattern = str(primary.get("elasticity_pattern", "Not assessable"))
    quadrant = {
        "Price ↑ / Qty ↓": "PRICE_UP_QTY_DOWN",
        "Price ↓ / Qty ↑": "PRICE_DOWN_QTY_UP",
        "Price ↑ / Qty ↑": "PRICE_UP_QTY_UP",
        "Price ↓ / Qty ↓": "PRICE_DOWN_QTY_DOWN",
    }.get(pattern, "STABLE_OR_NOT_ASSESSABLE")
    competitor_rows = []
    for brand in ("BRAUN", "PANASONIC", "FLYCO"):
        item = competitors.get(brand)
        gap = relative.get(brand)
        competitor_rows.append({
            "brand": brand,
            "status": "OBSERVED" if item and item.get("valid_months", 0) >= 2 else "NOT_ASSESSABLE",
            "observation_id": item.get("observation_id") if item else None,
            "model": item.get("model") if item else None,
            "asp_lm_change_pct": item.get("asp_lm_change_pct") if item else None,
            "qty_lm_change_pct": item.get("qty_lm_change_pct") if item else None,
            "gmv_lm_change_pct": item.get("gmv_lm_change_pct") if item else None,
            "ph_median_price_premium_pct": gap.get("ph_price_premium_pct") if gap else None,
        })
    claims = package.claims
    return {
        "framework_version": "pricing-ai-analysis-v1",
        "level_1_descriptive_analysis": {
            "observation_id": primary.get("observation_id"),
            "asp": {"latest": primary.get("latest_asp"), "lm_change_pct": primary.get("asp_lm_change_pct"), "period_change_pct": primary.get("price_change_pct"), "min": primary.get("min_asp"), "max": primary.get("max_asp")},
            "qty": {"latest": primary.get("latest_qty"), "lm_change_pct": primary.get("qty_lm_change_pct"), "period_change_pct": primary.get("sales_change_pct"), "min": primary.get("min_qty"), "max": primary.get("max_qty")},
            "gmv": {"latest": primary.get("latest_gmv"), "lm_change_pct": primary.get("gmv_lm_change_pct"), "period_change_pct": primary.get("gmv_change_pct")},
        },
        "level_2_price_elasticity": {"quadrant": quadrant, "display_label": pattern, "pearson": primary.get("pearson"), "spearman": primary.get("spearman"), "price_leads_qty_1m": primary.get("price_leads_sales_1m"), "anomaly": primary.get("anomaly"), "anomaly_threshold_pct": primary.get("anomaly_threshold_pct"), "material_price_moves": primary.get("material_price_moves", [])},
        "level_3_competitor_analysis": competitor_rows,
        "monthly_report": {
            "executive_summary": {"market_change_claim_ids": [claims[0].claim_id, claims[6].claim_id], "key_risk_claim_id": claims[7].claim_id, "key_opportunity_claim_id": claims[8].claim_id},
            "pricing_performance": {"asp_top_down_observation_id": primary.get("observation_id"), "category_trend_scope": "selected Philips SKU plus validated mapped competitor models; not a full-market category total"},
            "competitor_watch": {"price_movement_claim_ids": [item.claim_id for item in package.external_drivers[:3]], "promotion_evidence_claim_id": package.external_drivers[3].claim_id, "official_evidence_count": len(evidence)},
            "ai_generated_insights": {"top_10_claim_ids": [item.claim_id for item in claims], "decision_claim_id": package.decision_claim.claim_id},
        },
    }


class OfficialEvidenceProvider:
    def __init__(self, registry_path: Path) -> None:
        self.endpoint = os.getenv("SEARXNG_URL", "").rstrip("/"); self.registry = json.loads(registry_path.read_text(encoding="utf-8")) if registry_path.exists() else {}

    def search(self, brand: str, model: str, periods: list[str]) -> list[EvidenceRecord]:
        domains = self.registry.get(brand.upper(), {}).get("domains", [])
        if not self.endpoint or not domains or not periods: return []
        query = f"({' OR '.join(f'site:{domain}' for domain in domains)}) {brand} {model} {' OR '.join(periods)}"
        try:
            response = httpx.get(f"{self.endpoint}/search", params={"q": query, "format": "json", "language": "zh-CN", "safesearch": 1}, timeout=30); response.raise_for_status(); results = response.json().get("results", [])[:10]
        except Exception: return []
        now = datetime.now(timezone.utc).isoformat(); evidence = []
        for result in results:
            url = str(result.get("url", "")); domain = next((item for item in domains if item.lower() in url.lower()), "")
            if not domain: continue
            snippet = str(result.get("content", ""))[:800]
            evidence.append(EvidenceRecord(f"EVENT-{len(evidence)+1:03d}", brand, model, periods[-1], str(result.get("title", "")), url, result.get("publishedDate"), now, domain, True, snippet, hashlib.sha256(f"{url}\n{snippet}".encode()).hexdigest()))
        return evidence


def deterministic_claims(observations: list[dict[str, Any]], evidence: list[EvidenceRecord]) -> list[InsightClaim]:
    primary = next((item for item in observations if item.get("role") == "PH"), {}); competitors = [item for item in observations if item.get("role") == "Competitor"]; primary_id = primary.get("observation_id", ""); primary_ids = [primary_id] if primary_id else []
    comp_by_brand = {item.get("brand", "").upper(): item for item in competitors}; focus = [brand for brand in ("BRAUN", "PANASONIC", "FLYCO") if brand in comp_by_brand]; external_ids = [item.evidence_id for item in evidence]
    def value(item: dict[str, Any], field: str) -> str:
        current = item.get(field); return "N/A" if current is None else f"{current:+.1f}%"
    if primary.get("valid_months", 0) < 2:
        descriptive, level = "Philips 有效月度数据不足，无法输出 ASP、Qty 与 GMV 的月度变化。", "NOT_ASSESSABLE"
    else:
        descriptive = f"Philips 最新月 ASP 环比{value(primary, 'asp_lm_change_pct')}，Qty 环比{value(primary, 'qty_lm_change_pct')}，GMV 环比{value(primary, 'gmv_lm_change_pct')}。"; level = "OBSERVED"
    elasticity = f"Philips 最新月弹性象限为 {primary.get('elasticity_pattern', 'Not assessable')}；同期相关系数仅描述共变，不单独证明价格导致销量变化。"
    claims = [
        InsightClaim(claim_id="CLAIM-001", claim_type="internal", subject="PHILIPS", text=descriptive, observation_ids=primary_ids, evidence_level=level, missing_data=[] if level == "OBSERVED" else ["ASP", "Qty", "GMV"]),
        InsightClaim(claim_id="CLAIM-002", claim_type="internal", subject="PHILIPS", text=elasticity, observation_ids=primary_ids, evidence_level="OBSERVED" if primary.get("pearson") is not None else "NOT_ASSESSABLE", alternative_explanations=["promotion", "traffic", "inventory"]),
    ]
    for brand in ("BRAUN", "PANASONIC", "FLYCO"):
        item = comp_by_brand.get(brand)
        if item:
            text = f"{brand} {item.get('model')} 最新月 ASP 环比{value(item, 'asp_lm_change_pct')}，Qty 环比{value(item, 'qty_lm_change_pct')}，GMV 环比{value(item, 'gmv_lm_change_pct')}，弹性象限为 {item.get('elasticity_pattern')}。"; ids = [item["observation_id"]]; evidence_level = "OBSERVED"
        else:
            text = f"{brand} 未找到可发布的对应型号或有效月度数据。"; ids = []; evidence_level = "NOT_ASSESSABLE"
        claims.append(InsightClaim(claim_id=f"CLAIM-{len(claims)+1:03d}", claim_type="external", subject=brand, text=text, observation_ids=ids, evidence_level=evidence_level, missing_data=[] if ids else [f"{brand} mapping or metrics"]))
    price_positions = [item for item in observations if item.get("type") == "relative_price" and str(item.get("brand", "")).upper() in {"BRAUN", "PANASONIC", "FLYCO"} and item.get("ph_price_premium_pct") is not None]
    if price_positions:
        nearest = min(price_positions, key=lambda item: abs(float(item["ph_price_premium_pct"])))
        position = f"Philips 中位 ASP 相对映射的 {nearest['brand']} {nearest['model']} 差异为{value(nearest, 'ph_price_premium_pct')}；该比较仅覆盖映射型号，不代表完整品类。"
    else: position = "当前映射数据不足，无法稳定判断 Philips 的竞品价格位置。"
    claims.append(InsightClaim(claim_id="CLAIM-006", claim_type="integrated", subject="INTEGRATED", text=position, observation_ids=[nearest["observation_id"]] if price_positions else [], evidence_level="OBSERVED" if price_positions else "NOT_ASSESSABLE", missing_data=[] if price_positions else ["competitor metrics"]))
    event_text = "已检索到官方同期信息，可作为外因线索；仍需结合促销、流量和库存数据验证时间对应与因果。" if evidence else "未检索到可验证的官方同期事件；当前不得把销量变化归因于竞品决策、促销或其他外因。"
    claims.append(InsightClaim(claim_id="CLAIM-007", claim_type="external", subject="OFFICIAL_EVENTS", text=event_text, observation_ids=primary_ids, external_evidence_ids=external_ids, evidence_level="SUPPORTED" if evidence else "HYPOTHESIS", alternative_explanations=["promotion", "traffic", "inventory"], missing_data=[] if evidence else ["verified official event evidence"]))
    anomaly_text = "最新月 Qty 变化超过历史波动阈值，已标记为异常 SKU 月份，需优先核查促销、断货、流量与口径变化。" if primary.get("anomaly") else "最新月 Qty 变化未超过当前历史波动阈值，但仍应结合促销、库存和流量监测。"
    claims.append(InsightClaim(claim_id="CLAIM-008", claim_type="integrated", subject="INTEGRATED", text=anomaly_text, observation_ids=primary_ids, evidence_level="OBSERVED", missing_data=["promotion", "traffic", "inventory"])); claims.append(InsightClaim(claim_id="CLAIM-009", claim_type="integrated", subject="INTEGRATED", text="机会：围绕有销量改善且 GMV/毛利可守住的价格区间做小范围验证；风险：仅凭相关性降价可能侵蚀 GMV 与毛利。", observation_ids=primary_ids, evidence_level="HYPOTHESIS", missing_data=["gross margin"])); claims.append(InsightClaim(claim_id="CLAIM-010", claim_type="decision", subject="DECISION", text="下一步补齐促销、流量、库存和毛利数据，再做限定平台与周期的受控价格测试；以 Qty、GMV、毛利共同设门槛，未达标即恢复原方案。", observation_ids=primary_ids, external_evidence_ids=external_ids, evidence_level="HYPOTHESIS", missing_data=["promotion", "traffic", "inventory", "gross margin guardrail"])); return claims


def deterministic_decision(facts: dict[str, Any]) -> DecisionRecommendation:
    primary = next((item for item in facts.get("observations", []) if item.get("role") == "PH"), {})
    observation_ids = [primary["observation_id"]] if primary.get("observation_id") else []
    evidence_ids = [item["evidence_id"] for item in facts.get("external_evidence", [])]
    return DecisionRecommendation(
        action="先补齐促销、流量、库存与毛利信息，再在当前渠道开展可回滚的小范围价格验证。",
        channel=str(facts.get("channel", "JD")).upper(),
        timing="完成缺失数据核查并确认当前映射有效后",
        proposed_range="未获得毛利和促销约束前不设定具体价格区间",
        target_metrics=["Qty", "GMV", "Gross Margin"],
        guardrails=["不得以销量增长替代 GMV 与毛利检查", "竞品映射或库存状态改变时停止比较"],
        stop_condition="GMV 或毛利不满足业务门槛，或发现促销、流量、库存口径不可比时停止并回滚。",
        validation_design="在同一渠道与可比月份进行受控价格测试，同时记录 ASP、Qty、GMV、毛利、促销、流量与库存。",
        evidence_level="HYPOTHESIS",
        observation_ids=observation_ids,
        external_evidence_ids=evidence_ids,
        missing_data=["promotion", "traffic", "inventory", "gross margin"],
    )


def build_standard_facts(records: list[MonthlyMetric], channel: str, category: str = "Shaver") -> list[dict[str, Any]]:
    facts: list[dict[str, Any]] = []
    groups: dict[tuple[str, str], list[MonthlyMetric]] = {}
    for record in records: groups.setdefault((record.brand, record.model), []).append(record)
    for (brand, model), items in groups.items():
        ordered = sorted(items, key=lambda item: item.period); by_period = {(item.period.year, item.period.month): item for item in ordered}
        for index, item in enumerate(ordered):
            previous = ordered[index - 1] if index else None; prior_year = by_period.get((item.period.year - 1, item.period.month))
            facts.append({"month": item.period.isoformat(), "brand": brand, "sku": model, "category": category, "platform": channel, "asp": item.price, "qty": item.sales, "gmv": item.sales_value, "asp_lm_change_pct": percent_change(item.price, previous.price) if previous else None, "qty_lm_change_pct": percent_change(item.sales, previous.sales) if previous else None, "gmv_lm_change_pct": percent_change(item.sales_value, previous.sales_value) if previous else None, "asp_ly_change_pct": percent_change(item.price, prior_year.price) if prior_year else None, "qty_ly_change_pct": percent_change(item.sales, prior_year.sales) if prior_year else None, "gmv_ly_change_pct": percent_change(item.sales_value, prior_year.sales_value) if prior_year else None, "mapping_state": item.mapping_state, "source_sheet": item.source_sheet, "source_cells": {key: asdict(value) for key, value in item.source_cells.items()}})
    return sorted(facts, key=lambda item: (item["month"], item["brand"], item["sku"]))


class DuckDBRepository:
    def __init__(self, path: Path) -> None: self.path = path; path.parent.mkdir(parents=True, exist_ok=True)

    def save(self, run_id: str, channel: str, sku: str, source_hash: str, mappings: list[MappingItem], records: list[MonthlyMetric], evidence: list[EvidenceRecord], claims: list[InsightClaim], quality: dict[str, Any], standard_facts: list[dict[str, Any]] | None = None, decision: DecisionRecommendation | None = None) -> None:
        con = duckdb.connect(str(self.path))
        try:
            con.execute("BEGIN TRANSACTION")
            con.execute("CREATE TABLE analysis_run(run_id VARCHAR, channel VARCHAR, primary_sku VARCHAR, source_sha256 VARCHAR, status VARCHAR, generated_at TIMESTAMP, quality_json JSON)")
            con.execute("CREATE TABLE sku_mapping(run_id VARCHAR, brand VARCHAR, model VARCHAR, role VARCHAR, source VARCHAR, confidence DOUBLE, mapping_state VARCHAR, raw_mapping_value VARCHAR)")
            con.execute("CREATE TABLE monthly_metric(run_id VARCHAR, period DATE, period_label VARCHAR, brand VARCHAR, model VARCHAR, role VARCHAR, price DOUBLE, sales DOUBLE, sales_value DOUBLE, source_sheet VARCHAR, source_cells_json JSON, mapping_state VARCHAR)")
            con.execute("CREATE TABLE standard_fact(run_id VARCHAR, month DATE, brand VARCHAR, sku VARCHAR, category VARCHAR, platform VARCHAR, asp DOUBLE, qty DOUBLE, gmv DOUBLE, asp_lm_change_pct DOUBLE, qty_lm_change_pct DOUBLE, gmv_lm_change_pct DOUBLE, asp_ly_change_pct DOUBLE, qty_ly_change_pct DOUBLE, gmv_ly_change_pct DOUBLE, mapping_state VARCHAR, source_sheet VARCHAR, source_cells_json JSON)")
            con.execute("CREATE TABLE external_evidence(run_id VARCHAR, evidence_id VARCHAR, brand VARCHAR, model VARCHAR, period_label VARCHAR, title VARCHAR, url VARCHAR, published_at VARCHAR, retrieved_at VARCHAR, source_domain VARCHAR, official BOOLEAN, snippet VARCHAR, content_sha256 VARCHAR)")
            con.execute("CREATE TABLE insight_claim(run_id VARCHAR, claim_id VARCHAR, claim_type VARCHAR, subject VARCHAR, claim VARCHAR, observation_ids_json JSON, external_evidence_ids_json JSON, evidence_level VARCHAR, alternatives_json JSON, missing_data_json JSON)")
            con.execute("CREATE TABLE decision_recommendation(run_id VARCHAR, decision_json JSON)")
            con.execute("INSERT INTO analysis_run VALUES (?, ?, ?, ?, ?, ?, ?)", [run_id, channel, sku, source_hash, "completed", datetime.now(timezone.utc).replace(tzinfo=None), json.dumps(quality)])
            if mappings: con.executemany("INSERT INTO sku_mapping VALUES (?, ?, ?, ?, ?, ?, ?, ?)", [[run_id, item.brand, item.model, item.role, item.source, item.confidence, item.mapping_state, item.raw_mapping_value] for item in mappings])
            if records: con.executemany("INSERT INTO monthly_metric VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)", [[run_id, item.period, item.period_label, item.brand, item.model, item.role, item.price, item.sales, item.sales_value, item.source_sheet, json.dumps({key: asdict(value) for key, value in item.source_cells.items()}, ensure_ascii=False), item.mapping_state] for item in records])
            if standard_facts: con.executemany("INSERT INTO standard_fact VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)", [[run_id, item["month"], item["brand"], item["sku"], item["category"], item["platform"], item["asp"], item["qty"], item["gmv"], item["asp_lm_change_pct"], item["qty_lm_change_pct"], item["gmv_lm_change_pct"], item["asp_ly_change_pct"], item["qty_ly_change_pct"], item["gmv_ly_change_pct"], item["mapping_state"], item["source_sheet"], json.dumps(item["source_cells"], ensure_ascii=False)] for item in standard_facts])
            if evidence: con.executemany("INSERT INTO external_evidence VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)", [[run_id, *asdict(item).values()] for item in evidence])
            if claims: con.executemany("INSERT INTO insight_claim VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)", [[run_id, item.claim_id, item.claim_type, item.subject, item.text, json.dumps(item.observation_ids), json.dumps(item.external_evidence_ids), item.evidence_level, json.dumps(item.alternative_explanations), json.dumps(item.missing_data)] for item in claims])
            if decision: con.execute("INSERT INTO decision_recommendation VALUES (?, ?)", [run_id, decision.model_dump_json()])
            con.execute("COMMIT")
            if con.execute("SELECT count(*) FROM monthly_metric WHERE run_id=?", [run_id]).fetchone()[0] != len(records): raise RuntimeError("DuckDB read-back row count mismatch")
            if standard_facts and con.execute("SELECT count(*) FROM standard_fact WHERE run_id=?", [run_id]).fetchone()[0] != len(standard_facts): raise RuntimeError("DuckDB standard fact read-back row count mismatch")
        except Exception:
            try: con.execute("ROLLBACK")
            except Exception: pass
            raise
        finally: con.close()


def compose_pdf(preview_dir: Path, output_pdf: Path, safe_sku: str) -> None:
    first = preview_dir / f"ppt-{safe_sku}.png"; extra = sorted(preview_dir.glob(f"ppt-{safe_sku}-slide*.png"), key=lambda path: int(re.search(r"slide(\d+)", path.stem).group(1))); slides = [first, *extra]
    if not slides or not all(path.exists() for path in slides): raise FileNotFoundError("PPT previews required for PDF composition")
    pdf = canvas.Canvas(str(output_pdf), pagesize=(960, 1080), pageCompression=1); pdf.setTitle(f"SKU Data Insight - {safe_sku}")
    for index in range(0, len(slides), 2):
        pair = slides[index:index + 2]; pdf.drawImage(ImageReader(str(pair[0])), 0, 540, width=960, height=540, preserveAspectRatio=True, mask="auto")
        if len(pair) == 2: pdf.drawImage(ImageReader(str(pair[1])), 0, 0, width=960, height=540, preserveAspectRatio=True, mask="auto")
        pdf.showPage()
    pdf.save()


def create_employee_archive(pdf_path: Path, json_path: Path, zip_path: Path, safe_sku: str) -> None:
    """Package only employee-facing deliverables; validation artifacts stay internal."""
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.write(pdf_path, arcname=pdf_path.name)
        archive.write(json_path, arcname=f"SKU_Data_Insight_{safe_sku}.json")


class InsightPipeline:
    def __init__(self, root: Path, template_path: Path) -> None:
        self.root = root
        self.template_path = template_path
        self.contract_path = root / "data_contract.yaml"
        self.runs_dir = root / "outputs" / "agent-runs" / ".staging"
        self.runs_dir.mkdir(parents=True, exist_ok=True)
        credential_source = load_agnes_api_key(root)
        if credential_source and "AGNES_ALLOW_FALLBACK" not in os.environ:
            os.environ["AGNES_ALLOW_FALLBACK"] = "false"
        self.agnes = AgnesClient(root / ".work" / "llm-audit.jsonl", credential_source)
        self.evidence_provider = OfficialEvidenceProvider(root / "official_sources.json")
        self.semantic_cache = SemanticPlanCache(root / ".work" / "semantic-plans")

    def _semantic_plan(self, adapter: WorkbookAdapter, *, labels_confirmed: bool = False) -> tuple[WorkbookSemanticPlan, WorkbookStructureManifest]:
        profiler = adapter.semantic_profiler()
        manifest = profiler.build_manifest()
        prompt_version = str(adapter.config.data.get("semantic", {}).get("prompt_version", "semantic-plan-v1"))
        recognizer_model = self.agnes.model if self.agnes.enabled else None
        cached = self.semantic_cache.read(manifest.layout_fingerprint, prompt_version=prompt_version, recognizer_model=recognizer_model)
        if cached is not None:
            validated = profiler.validate_plan(cached, manifest)
            if labels_confirmed and validated.review_required:
                validated = validated.model_copy(update={"review_required": False, "human_approved": True, "review_reasons": [*validated.review_reasons, "Approved by employee for this structural fingerprint"]})
                self.semantic_cache.write(validated)
            elif cached.human_approved or not cached.review_required:
                validated = validated.model_copy(update={"review_required": False, "human_approved": cached.human_approved, "review_reasons": cached.review_reasons})
            return validated, manifest
        try:
            draft = profiler.deterministic_plan(manifest)
        except Exception:
            draft = None
        candidate = self.agnes.recognize_semantic_plan(manifest, draft)
        candidate = candidate.model_copy(update={"prompt_version": prompt_version, "recognizer_model": recognizer_model})
        plan = profiler.validate_plan(candidate, manifest)
        if labels_confirmed and plan.review_required:
            plan = plan.model_copy(update={"review_required": False, "human_approved": True, "review_reasons": [*plan.review_reasons, "Approved by employee for this structural fingerprint"]})
        self.semantic_cache.write(plan)
        return plan, manifest

    def prepare(
        self,
        workbook_path: Path,
        intent: AnalysisIntent,
        confirmed: dict[str, list[ConfirmedMapping]] | None = None,
        *,
        semantic_labels_confirmed: bool = False,
    ) -> dict[str, Any]:
        adapter = WorkbookAdapter(workbook_path, self.contract_path)
        try:
            semantic_plan, structure_manifest = self._semantic_plan(adapter, labels_confirmed=semantic_labels_confirmed)
            mappings_by_channel: dict[str, list[MappingItem]] = {}
            resolutions_by_channel: dict[str, list[MappingResolution]] = {}
            pending: dict[str, list[dict[str, Any]]] = {}
            target_skus: dict[str, str] = {}
            for target in intent.analysis_targets:
                channel = target.channel
                target_skus[channel] = target.primary_sku
                approved = {item.brand: item.model for item in (confirmed or {}).get(channel, [])}
                mappings, resolutions = adapter.resolve_mapping(
                    channel,
                    target.primary_sku,
                    approved,
                    semantic_plan=semantic_plan,
                    months=intent.months,
                )
                inventory = adapter.model_inventory(channel)
                for brand, selected in approved.items():
                    if any(item.brand == brand for item in mappings):
                        continue
                    if not any(normalize(model) == normalize(selected) for model in inventory.get(brand, [])):
                        raise WorkbookContractError("CONFIRMED_MODEL_NOT_FOUND", f"Confirmed model {brand} {selected} is absent from {channel} Raw Data")
                    raise WorkbookContractError("CONFIRMED_MODEL_NOT_DECLARED", f"Confirmed model {brand} {selected} is not declared for {target.primary_sku} in {channel}")
                mappings_by_channel[channel] = mappings
                resolutions_by_channel[channel] = resolutions
                pending_rows: list[dict[str, Any]] = []
                for resolution in resolutions:
                    if resolution.state not in {"MULTIPLE_CANDIDATES", "LIFECYCLE_SPLIT_REQUIRED"} or resolution.brand in approved:
                        continue
                    pending_rows.extend(
                        {
                            "brand": resolution.brand,
                            "model": model,
                            "confidence": 1.0,
                            "coverage": resolution.coverage.get(model),
                            "available_periods": resolution.available_periods.get(model, []),
                            "median_asp": None,
                            "price_distance": None,
                            "reasons": [resolution.reason],
                        }
                        for model in resolution.present_candidates
                    )
                if pending_rows:
                    pending[channel] = pending_rows
            return {
                "mappings": mappings_by_channel,
                "resolutions": resolutions_by_channel,
                "pending": pending,
                "adapter_version": adapter.config.adapter_version,
                "semantic_plan": semantic_plan,
                "structure_manifest": structure_manifest,
                "target_skus": target_skus,
            }
        finally:
            adapter.close()

    def run(
        self,
        workbook_path: Path,
        intent: AnalysisIntent,
        confirmed: dict[str, list[ConfirmedMapping]] | None = None,
        *,
        semantic_labels_confirmed: bool = False,
    ) -> dict[str, Any]:
        prepared = self.prepare(workbook_path, intent, confirmed, semantic_labels_confirmed=semantic_labels_confirmed)
        semantic_plan: WorkbookSemanticPlan = prepared["semantic_plan"]
        if semantic_plan.review_required and not semantic_labels_confirmed:
            return {
                "status": "awaiting_label_confirmation",
                "intent": intent.model_dump(),
                "semantic_plan": semantic_plan.model_dump(),
            }
        if prepared["pending"]:
            return {
                "status": "awaiting_mapping_confirmation",
                "intent": intent.model_dump(),
                "mapping_candidates": prepared["pending"],
                "mapping_resolutions": {channel: serialize_resolutions(items) for channel, items in prepared["resolutions"].items()},
                "semantic_plan": semantic_plan.model_dump(),
            }
        source_hash = sha256_file(workbook_path)
        run_group = uuid.uuid4().hex
        temp_group = self.runs_dir / f"{run_group}.tmp"
        final_group = self.root / "outputs" / "agent-runs" / run_group
        if temp_group.exists():
            shutil.rmtree(temp_group)
        temp_group.mkdir(parents=True)
        adapter = WorkbookAdapter(workbook_path, self.contract_path)
        channel_results: list[dict[str, Any]] = []
        try:
            for target in intent.analysis_targets:
                channel = target.channel
                sku = target.primary_sku
                channel_dir = temp_group / channel
                channel_dir.mkdir()
                channel_run_id = f"{run_group}-{channel.lower()}"
                mappings = prepared["mappings"][channel]
                records = adapter.extract(channel, mappings, intent.months)
                quality = validate_metrics(records, sku, intent.months, adapter.config.data["quality"])
                observations = build_observations(records, mappings)
                standard_facts = build_standard_facts(records, channel)
                periods = sorted({item.period_label for item in records})
                evidence: list[EvidenceRecord] = []
                for mapping in mappings:
                    if mapping.role == "Competitor":
                        evidence.extend(self.evidence_provider.search(mapping.brand, mapping.model, periods))
                unique_evidence: list[EvidenceRecord] = []
                seen_evidence: set[str] = set()
                for item in evidence:
                    if item.content_sha256 in seen_evidence:
                        continue
                    seen_evidence.add(item.content_sha256)
                    unique_evidence.append(replace(item, evidence_id=f"EVENT-{len(unique_evidence)+1:03d}"))
                evidence = unique_evidence
                fallback_claims = deterministic_claims(observations, evidence)
                facts = {
                    "channel": channel,
                    "primary_sku": sku,
                    "period_scope": periods,
                    "observations": observations,
                    "external_evidence": [asdict(item) for item in evidence],
                    "quality_warnings": quality["warnings"],
                    "missing_driver_inputs": ["promotion", "traffic", "inventory", "rating", "listing", "media", "gross margin"],
                    "method_constraints": [
                        "ASP/Qty association is descriptive and does not prove causality",
                        "external causes require allow-listed official evidence IDs aligned to the period",
                        "recommendations without causal evidence must be framed as controlled tests",
                    ],
                }
                fallback_package = InsightPackage(claims=fallback_claims, decision=deterministic_decision(facts))
                insight_package = self.agnes.synthesize_insights(facts, fallback_package)
                claims = insight_package.claims
                analysis_framework = build_pricing_analysis_framework(observations, insight_package, evidence)
                insight_agent_trace = [
                    {"step": "metric_diagnostics", "tool": "deterministic_analysis_engine", "status": "completed", "observation_count": len(observations)},
                    {"step": "official_evidence", "tool": "OfficialEvidenceProvider", "status": "completed" if evidence else "no_verified_evidence", "evidence_count": len(evidence)},
                    {"step": "insight_generation", "tool": "AgnesClient", "status": self.agnes.status().get("last_status"), "model": self.agnes.status().get("model")},
                    {"step": "policy_validation", "tool": "InsightPackageValidator", "status": "completed", "claim_count": len(claims)},
                ]
                safe_sku = re.sub(r"[^A-Za-z0-9_-]+", "-", sku)
                db_path = channel_dir / f"SKU_Data_Insight_{safe_sku}.duckdb"
                DuckDBRepository(db_path).save(channel_run_id, channel, sku, source_hash, mappings, records, evidence, claims, quality, standard_facts, insight_package.decision)
                analysis = {
                    "intent": intent.model_dump(),
                    "semantic_plan": semantic_plan.model_dump(),
                    "mapping": serialize_mappings(mappings),
                    "mapping_resolutions": serialize_resolutions(prepared["resolutions"][channel]),
                    "quality": quality,
                    "observations": observations,
                    "standard_facts": standard_facts,
                    "external_evidence": [asdict(item) for item in evidence],
                    "claims": [item.model_dump() for item in claims],
                    "decision": insight_package.decision.model_dump(),
                    "pricing_analysis_framework": analysis_framework,
                    "insight_agent_trace": insight_agent_trace,
                    "llm_status": self.agnes.status(),
                }
                write_json(channel_dir / "analysis.json", analysis)
                payload = {
                    "payload_version": RENDER_PAYLOAD_VERSION,
                    "run_id": channel_run_id,
                    "source_path": str(workbook_path),
                    "source_sha256": source_hash,
                    "source_sheet": adapter.config.channel_sheet(channel),
                    "adapter_version": prepared["adapter_version"],
                    "semantic_plan_fingerprint": semantic_plan.layout_fingerprint,
                    "database_path": str(db_path),
                    "sku": sku,
                    "safe_sku": safe_sku,
                    "channel": channel,
                    "periods": periods,
                    "mapping": serialize_mappings(mappings),
                    "records": [item.to_dict() for item in records],
                    "standard_facts": standard_facts,
                    "observations": observations,
                    "insights": [item.text for item in claims],
                    "claims": [item.model_dump() for item in claims],
                    "decision": insight_package.decision.model_dump(),
                    "pricing_analysis_framework": analysis_framework,
                    "insight_agent_trace": insight_agent_trace,
                    "quality": quality,
                }
                payload_path = channel_dir / "render_payload.json"
                write_json(payload_path, payload)
                self._render(payload_path, channel_dir)
                pdf_path = channel_dir / f"SKU_Data_Insight_{safe_sku}.pdf"
                compose_pdf(channel_dir / "previews", pdf_path, safe_sku)
                self._verify_artifacts(channel_dir, safe_sku, records)
                (channel_dir / f"SKU_Data_Insight_{safe_sku}.pptx").unlink()
                channel_results.append({
                    "channel": channel,
                    "sku": sku,
                    "safe_sku": safe_sku,
                    "quality": quality,
                    "claims": [item.model_dump() for item in claims],
                    "decision": insight_package.decision.model_dump(),
                    "mapping": serialize_mappings(mappings),
                    "records": [item.to_dict() for item in records],
                })
            write_json(temp_group / "run_manifest.json", self._build_manifest(temp_group, run_group, source_hash, intent, prepared["adapter_version"], semantic_plan))
            final_group.parent.mkdir(parents=True, exist_ok=True)
            if final_group.exists():
                raise RuntimeError("Run output directory already exists")
            self._atomic_publish(temp_group, final_group)
            results: list[dict[str, Any]] = []
            for item in channel_results:
                channel, safe_sku = item["channel"], item["safe_sku"]
                channel_dir = final_group / channel
                pdf_path = channel_dir / f"SKU_Data_Insight_{safe_sku}.pdf"
                json_path = channel_dir / "analysis.json"
                zip_path = final_group / f"SKU_Data_Insight_{safe_sku}_{channel}.zip"
                create_employee_archive(pdf_path, json_path, zip_path, safe_sku)
                results.append({
                    "channel": channel,
                    "sku": item["sku"],
                    "output_dir": str(channel_dir),
                    "quality": item["quality"],
                    "claims": item["claims"],
                    "decision": item["decision"],
                    "mapping": item["mapping"],
                    "records": item["records"],
                    "pdf": str(pdf_path),
                    "json": str(json_path),
                    "manifest": str(final_group / "run_manifest.json"),
                    "zip": str(zip_path),
                })
            return {"status": "completed", "run_id": run_group, "intent": intent.model_dump(), "semantic_plan": semantic_plan.model_dump(), "llm_status": self.agnes.status(), "results": results}
        except Exception as exc:
            if not isinstance(exc, WorkbookContractError) or exc.code != "ATOMIC_PUBLISH_LOCKED":
                shutil.rmtree(temp_group, ignore_errors=True)
            raise
        finally:
            adapter.close()

    def _render(self, payload_path: Path, output_dir: Path) -> None:
        node = os.getenv("RUNTIME_NODE", r"C:\Users\320332974\.cache\codex-runtimes\codex-primary-runtime\dependencies\node\bin\node.exe"); env = os.environ.copy(); env.setdefault("RUNTIME_PYTHON", r"C:\Users\320332974\.cache\codex-runtimes\codex-primary-runtime\dependencies\python\python.exe"); env.setdefault("RUNTIME_NODE_MODULES", r"C:\Users\320332974\.cache\codex-runtimes\codex-primary-runtime\dependencies\node\node_modules")
        process = subprocess.run([node, str(self.root / "sku_insight_agent.mjs"), "--payload", str(payload_path), "--template", str(self.template_path), "--output", str(output_dir)], cwd=self.root, env=env, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=300)
        if process.returncode != 0: raise RuntimeError(f"Renderer failed: {process.stderr or process.stdout}")

    @staticmethod
    def _atomic_publish(temp_group: Path, final_group: Path) -> None:
        last_error: PermissionError | None = None
        for delay in (0.0, 0.25, 0.5, 1.0, 2.0, 4.0):
            if delay: time.sleep(delay)
            try:
                temp_group.replace(final_group); return
            except PermissionError as exc: last_error = exc
        raise WorkbookContractError("ATOMIC_PUBLISH_LOCKED", "Validated artifacts could not be atomically published because OneDrive or another process is holding the staging directory. The validated staging output was preserved; retry after synchronization finishes.", {"staging_dir": str(temp_group), "target_dir": str(final_group), "error": str(last_error)})

    @staticmethod
    def _verify_artifacts(channel_dir: Path, safe_sku: str, records: list[MonthlyMetric]) -> None:
        xlsx = channel_dir / f"SKU_Data_Insight_{safe_sku}.xlsx"; pptx = channel_dir / f"SKU_Data_Insight_{safe_sku}.pptx"; pdf = channel_dir / f"SKU_Data_Insight_{safe_sku}.pdf"
        for path in (xlsx, pptx):
            if not path.exists() or path.stat().st_size < 1000 or not zipfile.is_zipfile(path): raise RuntimeError(f"Invalid generated artifact: {path.name}")
        if not pdf.exists() or pdf.stat().st_size < 1000 or pdf.read_bytes()[:4] != b"%PDF": raise RuntimeError(f"Invalid generated artifact: {pdf.name}")
        with zipfile.ZipFile(pptx) as archive:
            slide_xml = [name for name in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)]
            if len(slide_xml) != 3: raise RuntimeError(f"Generated monitoring report must contain 3 slides, found {len(slide_xml)}")
        with zipfile.ZipFile(xlsx) as archive:
            if not any(name.startswith("xl/media/") for name in archive.namelist()): raise RuntimeError("Generated Excel Dashboard is missing the embedded SKU trend chart")
        if len(re.findall(rb"/Type\s*/Page\b", pdf.read_bytes())) != 2: raise RuntimeError("Generated combined PDF must contain two pages")
        book = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
        try:
            required = {"Monthly Data", "Standard Fact Table", "Competitor Mapping", "Price Elasticity", "Competitor Tracking", "Monitoring Report"}
            if missing := sorted(required - set(book.sheetnames)): raise RuntimeError(f"Generated Excel is missing sheets: {missing}")
            rows = list(book["Monthly Data"].iter_rows(min_row=2, values_only=True))
            if len(rows) != len(records): raise RuntimeError("Generated Excel row count does not match canonical records")
            expected = [(item.period_label, item.brand, item.model, item.price, item.sales, item.sales_value) for item in records]
            actual = [(str(row[2]), str(row[4]), str(row[5]), float(row[6]) if row[6] is not None else None, float(row[7]) if row[7] is not None else None, float(row[8]) if row[8] is not None else None) for row in rows]
            if actual != expected: raise RuntimeError("Generated Excel values do not match canonical records")
            if book["Competitor Mapping"].iter_rows(min_row=1, max_row=1, values_only=True).__next__()[:4] != ("Philips SKU", "Competitor Brand", "Competitor SKU", "Mapping Type"): raise RuntimeError("Competitor Mapping leading columns do not match the output contract")
            if len(list(book["Standard Fact Table"].iter_rows(min_row=2, values_only=True))) != len(records): raise RuntimeError("Standard Fact Table row count does not match canonical records")
            if len(list(book["Monitoring Report"].iter_rows(min_row=4, values_only=True))) < 10: raise RuntimeError("Monitoring Report does not contain Top 10 insights")
        finally: book.close()

    def _build_manifest(self, group_dir: Path, run_id: str, source_hash: str, intent: AnalysisIntent, adapter_version: str, semantic_plan: WorkbookSemanticPlan) -> dict[str, Any]:
        artifacts = [{"path": str(path.relative_to(group_dir)), "bytes": path.stat().st_size, "sha256": sha256_file(path)} for path in sorted(item for item in group_dir.rglob("*") if item.is_file())]
        return {"run_id": run_id, "generated_at": datetime.now(timezone.utc).isoformat(), "source_sha256": source_hash, "intent": intent.model_dump(), "contract_version": "1.2", "adapter_version": adapter_version, "semantic_plan": semantic_plan.model_dump(), "llm_status": self.agnes.status(), "renderer_version": "v20-single-payload", "agnes_model": self.agnes.model if self.agnes.enabled else "local-fallback", "artifacts": artifacts}
