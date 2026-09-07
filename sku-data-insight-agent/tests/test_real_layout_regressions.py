from __future__ import annotations

from datetime import date
from pathlib import Path

from adapters import WorkbookAdapter
from tests.fixture_factory import make_fixture


CONTRACT = Path(__file__).parents[1] / "data_contract.yaml"


def test_newest_first_physical_blocks_do_not_cross_month_boundaries(tmp_path: Path) -> None:
    path = make_fixture(tmp_path / "raw.xlsx")
    adapter = WorkbookAdapter(path, CONTRACT)
    try:
        sheet = adapter.values_book[adapter.config.channel_sheet("JD")]
        blocks = adapter._month_blocks(sheet)
        assert blocks[0][1] == date(2025, 1, 1)
    finally:
        adapter.close()

    # Reverse complete monthly blocks to mirror the production workbook's
    # newest-first physical layout.
    import openpyxl

    book = openpyxl.load_workbook(path)
    sheet = book["3.2 Shaver SKU(JD)"]
    original = list(sheet.iter_rows(values_only=True))
    starts = [index for index, row in enumerate(original) if len(row) >= 4 and isinstance(row[3], str) and row[3].endswith(" Value")]
    chunks = []
    for position, start in enumerate(starts):
        end = starts[position + 1] if position + 1 < len(starts) else len(original)
        chunks.append(original[start:end])
    sheet.delete_rows(1, sheet.max_row)
    out_row = 1
    for chunk in reversed(chunks):
        for values in chunk:
            for col, value in enumerate(values, 1):
                sheet.cell(out_row, col, value)
            out_row += 1
    book.save(path)

    adapter = WorkbookAdapter(path, CONTRACT)
    try:
        mappings, _ = adapter.resolve_mapping("JD", "S3203/08")
        records = adapter.extract("JD", mappings, 12)
        assert len(records) == 60
        assert len({item.period for item in records if item.brand == "PHILIPS"}) == 12
    finally:
        adapter.close()
