from __future__ import annotations

from datetime import date
from pathlib import Path

import openpyxl


CHANNEL_SHEETS = {
    "JD": "3.2 Shaver SKU(JD)",
    "ALI": "3.1 Shaver SKU(ALi)",
    "OFFLINE": "3.3 Shaver SKU(Offline)",
}
MONTHS = ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")


MAPPINGS = {
    "S3203/08": {"BRAUN": "5603", "FLYCO": "FS966/967/968", "PANASONIC": "ES-RM31", "YOOSE": "C1"},
    "S1115/02": {"BRAUN": "300S", "FLYCO": "FS903", "PANASONIC": "ES-RC30", "YOOSE": "MINI"},
}


def make_fixture(path: Path, *, omit_s1115_from_jd_mapping: bool = False) -> Path:
    book = openpyxl.Workbook()
    key = book.active
    key.title = "Key SKUs List"
    row = 1
    for channel in CHANNEL_SHEETS:
        key.cell(row, 1, channel)
        row += 1
        for column, brand in enumerate(("PHILIPS", "BRAUN", "FLYCO", "PANASONIC", "YOOSE"), 1):
            key.cell(row, column, brand)
        row += 1
        for sku, competitors in MAPPINGS.items():
            if omit_s1115_from_jd_mapping and channel == "JD" and sku == "S1115/02":
                continue
            key.cell(row, 1, sku)
            for column, brand in enumerate(("BRAUN", "FLYCO", "PANASONIC", "YOOSE"), 2):
                key.cell(row, column, competitors[brand])
            row += 1
        row += 2

    for channel_index, (channel, sheet_name) in enumerate(CHANNEL_SHEETS.items()):
        sheet = book.create_sheet(sheet_name)
        cursor = 1
        channel_factor = 1 + channel_index * 0.08
        for month_index, month in enumerate(MONTHS):
            sheet.cell(cursor, 4, f"{month}2025 Value")
            cursor += 1
            for sku_index, sku in enumerate(MAPPINGS):
                price = round((319 - month_index * 2 - sku_index * 110) * channel_factor, 1)
                unit = float(120 + month_index * 11 + sku_index * 45 + (25 if month_index in (5, 10) else 0))
                sheet.cell(cursor, 2, sku)
                sheet.cell(cursor, 4, price * unit)
                sheet.cell(cursor, 7, unit)
                sheet.cell(cursor, 10, price)
                cursor += 1
            models_seen: set[tuple[str, str]] = set()
            for sku_index, competitors in enumerate(MAPPINGS.values()):
                for brand_index, (brand, model) in enumerate(competitors.items()):
                    key_pair = (brand, model)
                    if key_pair in models_seen:
                        continue
                    models_seen.add(key_pair)
                    price = round((285 - month_index * (1 + brand_index * 0.25) - sku_index * 90) * channel_factor, 1)
                    unit = float(105 + month_index * (8 + brand_index) + sku_index * 36)
                    sheet.cell(cursor, 13, model)
                    sheet.cell(cursor, 14, brand)
                    sheet.cell(cursor, 15, price * unit)
                    sheet.cell(cursor, 18, unit)
                    sheet.cell(cursor, 21, price)
                    cursor += 1
            cursor += 1

    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    return path
