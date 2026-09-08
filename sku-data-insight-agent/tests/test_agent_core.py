from __future__ import annotations

from pathlib import Path
import zipfile

import duckdb
import openpyxl
import pytest
from hypothesis import given, strategies as st

from adapters import WorkbookAdapter, WorkbookContractError, model_matches, primary_model_tokens
from pipeline import AgnesClient, AnalysisIntent, AnalysisPlan, ConfirmedMapping, DuckDBRepository, InsightPipeline, build_observations, build_standard_facts, create_employee_archive, default_analysis_plan, deterministic_claims, proto_l3_enabled, rule_based_intent, validate_analysis_plan, validate_metrics
from tests.fixture_factory import make_fixture


CONTRACT = Path(__file__).parents[1] / "data_contract.yaml"


def test_proto_l3_is_enabled_by_default(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.delenv("PROTO_L3_ENABLED", raising=False)
    assert proto_l3_enabled() is True


def test_invalid_llm_plan_falls_back_to_governed_plan(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    monkeypatch.setenv("AGNES_API_KEY", "test-key")
    monkeypatch.setenv("AGNES_ALLOW_FALLBACK", "false")
    client = AgnesClient(tmp_path / "audit.jsonl")
    intent = AnalysisIntent(primary_sku="S1115/02", channels=["JD"])
    invalid = default_analysis_plan(intent).model_copy(update={"steps": [default_analysis_plan(intent).steps[0], default_analysis_plan(intent).steps[7]]})
    monkeypatch.setattr(client, "_chat_model", lambda *args, **kwargs: invalid)
    result = client.generate_analysis_plan("分析", intent)
    assert result.source == "deterministic"
    assert result.steps[-1].operation == "render_report"


def test_governed_proto_l3_default_plan_is_valid() -> None:
    intent = AnalysisIntent(primary_sku="S1115/02", channels=["JD"])
    plan = validate_analysis_plan(default_analysis_plan(intent), intent)
    assert [step.operation for step in plan.steps] == [
        "inspect_workbook", "resolve_mapping", "extract_metrics", "validate_quality",
        "analyze_price_trend", "analyze_competitors", "search_official_evidence",
        "synthesize_insights", "generate_decision", "render_report",
    ]


def test_governed_proto_l3_rejects_cycles_and_skipped_gates() -> None:
    intent = AnalysisIntent(primary_sku="S1115/02", channels=["JD"])
    plan = default_analysis_plan(intent)
    cyclic = plan.model_copy(update={"steps": [plan.steps[0].model_copy(update={"depends_on": ["render"]}), *plan.steps[1:]]})
    with pytest.raises(ValueError, match="cycle"):
        validate_analysis_plan(cyclic, intent)
    skipped = plan.model_copy(update={"steps": [step for step in plan.steps if step.operation != "validate_quality"]})
    with pytest.raises(ValueError, match="skip"):
        validate_analysis_plan(skipped, intent)


def test_employee_archive_contains_only_pdf_and_json(tmp_path: Path) -> None:
    pdf_path = tmp_path / "SKU_Data_Insight_S1115-02.pdf"
    json_path = tmp_path / "analysis.json"
    zip_path = tmp_path / "employee.zip"
    pdf_path.write_bytes(b"%PDF-1.4\n")
    json_path.write_text('{"status":"completed"}', encoding="utf-8")

    create_employee_archive(pdf_path, json_path, zip_path, "S1115-02")

    with zipfile.ZipFile(zip_path) as archive:
        assert archive.namelist() == [
            "SKU_Data_Insight_S1115-02.pdf",
            "SKU_Data_Insight_S1115-02.json",
        ]


def test_intent_recognizes_sku_and_multiple_channels() -> None:
    intent = rule_based_intent("分析 Philips S1115/02 在京东、阿里和线下最近12个月的价格销量")
    assert intent.primary_sku == "S1115/02"
    assert intent.channels == ["JD", "ALI", "OFFLINE"]


def test_agent_request_requires_exactly_one_channel(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.delenv("AGNES_API_KEY", raising=False)
    client = AgnesClient()
    with pytest.raises(ValueError, match="一次只能分析一个销售渠道"):
        client.parse_intent("分析 Philips S1115/02 在 JD、ALI 和 Offline 最近12个月的价格销量")


def test_intent_preserves_same_channel_compound_philips_product_group() -> None:
    intent = rule_based_intent("分析 Philips YQ660/02/PQ663/02 在 ALI 最近 12 个月的价格与销量，并比较竞品")
    assert intent.channels == ["ALI"]
    assert intent.primary_sku == "YQ660/02 / PQ663/02"


def test_intent_preserves_shared_suffix_compound_group_next_to_channel() -> None:
    intent = rule_based_intent("分析 Philips X5005/X5001/X5002/X5003/00ALI最近12个月的价格与销量")
    assert intent.channels == ["ALI"]
    assert intent.primary_sku == "X5005 / X5001 / X5002 / X5003/00"


def test_intent_rejects_two_separate_philips_products() -> None:
    with pytest.raises(ValueError, match="一次只能分析一个 Philips 产品"):
        rule_based_intent("分析 Philips S3203/08 和 S1115/02 在 ALI 最近12个月")


def test_intent_ignores_series_descriptor_after_philips_model() -> None:
    intent = rule_based_intent("分析 Philips SP9888/63 / S9000 在 ALI 最近12个月")
    assert intent.channels == ["ALI"]
    assert intent.primary_sku == "SP9888/63"


def test_primary_tokens_support_series_words_and_omitted_prefixes() -> None:
    assert primary_model_tokens("SP9888/63 S9000 PRESTIGE") == {"SP988863"}
    assert primary_model_tokens("SERIES 6000 S6632/01") == {"S663201"}
    assert primary_model_tokens("XP9203/20 /9202/20 /9201/20 I9000 PRESTIGE") == {"XP920320", "XP920220", "XP920120"}


def test_model_matching_does_not_treat_product_name_prefix_as_an_alias() -> None:
    assert model_matches("MINI", "MINI")
    assert not model_matches("MINI-S", "MINI")
    assert not model_matches("MINI 2.0", "MINI")
    assert model_matches("C1 (ICE SHAVER)", "C1")
    assert model_matches("FS966/967/968", "FS967")


def test_mapping_is_channel_bounded_and_s1115_maps_fs903(tmp_path: Path) -> None:
    path = make_fixture(tmp_path / "raw.xlsx", omit_s1115_from_jd_mapping=True)
    adapter = WorkbookAdapter(path, CONTRACT)
    try:
        with pytest.raises(WorkbookContractError) as error:
            adapter.resolve_mapping("JD", "S1115/02")
        assert error.value.code == "PRIMARY_SKU_NOT_FOUND_IN_CHANNEL"
        assert error.value.details["channel"] == "JD"
        assert error.value.details["cross_channel_fallback"] is False
        ali_mapping, _ = adapter.resolve_mapping("ALI", "S1115/02")
        assert [(item.brand, item.model) for item in ali_mapping if item.brand == "FLYCO"] == [("FLYCO", "FS903")]
    finally:
        adapter.close()


def test_compound_philips_product_group_is_resolved_independently_in_all_channels(tmp_path: Path) -> None:
    path = make_fixture(tmp_path / "compound-groups.xlsx")
    book = openpyxl.load_workbook(path)
    key = book["Key SKUs List"]
    for row in range(1, key.max_row + 1):
        if key.cell(row, 1).value == "S3203/08":
            key.cell(row, 1, "PQ663/02")
    for sheet_name in ("3.1 Shaver SKU(ALi)", "3.2 Shaver SKU(JD)", "3.3 Shaver SKU(Offline)"):
        sheet = book[sheet_name]
        for row in range(1, sheet.max_row + 1):
            if sheet.cell(row, 2).value == "S3203/08":
                sheet.cell(row, 2, "YQ660/02 / PQ663/02")
    book.save(path)
    book.close()

    adapter = WorkbookAdapter(path, CONTRACT)
    try:
        for channel in ("JD", "ALI", "OFFLINE"):
            for requested in ("YQ660/02/PQ663/02", "YQ660/02", "PQ663/02"):
                mappings, _ = adapter.resolve_mapping(channel, requested)
                primary = next(item for item in mappings if item.role == "PH")
                assert primary.model == "YQ660/02 / PQ663/02"
                assert primary.mapping_state == "PRODUCT_GROUP_ALIAS"
            records = adapter.extract(channel, mappings, 12)
            assert len([item for item in records if item.role == "PH" and item.price is not None]) == 12
    finally:
        adapter.close()


def test_compound_product_group_does_not_borrow_mapping_from_another_channel(tmp_path: Path) -> None:
    path = make_fixture(tmp_path / "compound-channel-isolation.xlsx")
    book = openpyxl.load_workbook(path)
    key = book["Key SKUs List"]
    channel = None
    for row in range(1, key.max_row + 1):
        value = key.cell(row, 1).value
        if value in {"JD", "ALI", "OFFLINE"}:
            channel = value
        elif value == "S3203/08":
            key.cell(row, 1, None if channel == "OFFLINE" else "PQ663/02")
    for sheet_name in ("3.1 Shaver SKU(ALi)", "3.2 Shaver SKU(JD)", "3.3 Shaver SKU(Offline)"):
        sheet = book[sheet_name]
        for row in range(1, sheet.max_row + 1):
            if sheet.cell(row, 2).value == "S3203/08":
                sheet.cell(row, 2, "YQ660/02 / PQ663/02")
    book.save(path)
    book.close()

    adapter = WorkbookAdapter(path, CONTRACT)
    try:
        adapter.resolve_mapping("ALI", "YQ660/02/PQ663/02")
        with pytest.raises(WorkbookContractError) as error:
            adapter.resolve_mapping("OFFLINE", "YQ660/02/PQ663/02")
        assert error.value.code == "PRIMARY_SKU_NOT_FOUND_IN_CHANNEL"
        assert error.value.details["cross_channel_fallback"] is False
    finally:
        adapter.close()


def test_shared_suffix_product_group_matches_channel_mapping_representative(tmp_path: Path) -> None:
    path = make_fixture(tmp_path / "shared-suffix-group.xlsx")
    book = openpyxl.load_workbook(path)
    key = book["Key SKUs List"]
    channel = None
    for row in range(1, key.max_row + 1):
        value = key.cell(row, 1).value
        if value in {"JD", "ALI", "OFFLINE"}:
            channel = value
        elif channel == "ALI" and value == "S3203/08":
            key.cell(row, 1, "X5005/X5001/00")
    raw = book["3.1 Shaver SKU(ALi)"]
    for row in range(1, raw.max_row + 1):
        if raw.cell(row, 2).value == "S3203/08":
            raw.cell(row, 2, "X5005/X5001/X5002/X5003/00")
    book.save(path)
    book.close()

    adapter = WorkbookAdapter(path, CONTRACT)
    try:
        for requested in ("X5005/X5001/X5002/X5003/00", "X5002/00", "X5002"):
            mappings, _ = adapter.resolve_mapping("ALI", requested)
            primary = next(item for item in mappings if item.role == "PH")
            assert primary.model == "X5005/X5001/X5002/X5003/00"
            assert primary.raw_mapping_value == "X5005/X5001/00"
        records = adapter.extract("ALI", mappings, 12)
        assert len([item for item in records if item.role == "PH"]) == 12
    finally:
        adapter.close()


def test_series_descriptor_product_matches_suffixless_channel_key(tmp_path: Path) -> None:
    path = make_fixture(tmp_path / "series-descriptor.xlsx")
    book = openpyxl.load_workbook(path)
    key = book["Key SKUs List"]
    channel = None
    for row in range(1, key.max_row + 1):
        value = key.cell(row, 1).value
        if value in {"JD", "ALI", "OFFLINE"}:
            channel = value
        elif channel == "ALI" and value == "S3203/08":
            key.cell(row, 1, "SP9888")
    raw = book["3.1 Shaver SKU(ALi)"]
    month_index = -1
    for row in range(1, raw.max_row + 1):
        if str(raw.cell(row, 4).value or "").endswith(" Value"):
            month_index += 1
        if raw.cell(row, 2).value == "S3203/08":
            raw.cell(row, 2, "SP9888/63 S9000 PRESTIGE")
        elif raw.cell(row, 2).value == "S1115/02" and month_index < 6:
            raw.cell(row, 2, "SP9888")
    book.save(path)
    book.close()

    adapter = WorkbookAdapter(path, CONTRACT)
    try:
        for requested in ("SP9888/63 / S9000", "SP9888/63", "SP9888"):
            mappings, _ = adapter.resolve_mapping("ALI", requested)
            primary = next(item for item in mappings if item.role == "PH")
            assert primary.model == "SP9888/63 S9000 PRESTIGE"
            assert primary.raw_mapping_value == "SP9888"
        records = adapter.extract("ALI", mappings, 12)
        assert len([item for item in records if item.role == "PH"]) == 12
    finally:
        adapter.close()


def test_extracts_shape_provenance_and_validates(tmp_path: Path) -> None:
    path = make_fixture(tmp_path / "raw.xlsx")
    adapter = WorkbookAdapter(path, CONTRACT)
    try:
        mapping, _ = adapter.resolve_mapping("OFFLINE", "S3203/08")
        records = adapter.extract("OFFLINE", mapping, 12)
        assert len(records) == 60
        assert records[0].source_cells["price"].cell
        assert records[0].source_cells["sales"].sheet == "3.3 Shaver SKU(Offline)"
        quality = validate_metrics(records, "S3203/08", 12, adapter.config.data["quality"])
        assert quality["status"] == "passed"
    finally:
        adapter.close()


def test_missing_primary_mapping_fails_closed_without_competitor_inference(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    path = make_fixture(tmp_path / "raw.xlsx", omit_s1115_from_jd_mapping=True)
    monkeypatch.delenv("AGNES_API_KEY", raising=False)
    pipeline = InsightPipeline(Path(__file__).parents[1], tmp_path / "template.pptx")
    intent = AnalysisIntent(primary_sku="S1115/02", channels=["JD"])
    with pytest.raises(WorkbookContractError) as error:
        pipeline.prepare(path, intent)
    assert error.value.code == "PRIMARY_SKU_NOT_FOUND_IN_CHANNEL"


def test_empty_mapping_is_not_inferred(tmp_path: Path) -> None:
    path = make_fixture(tmp_path / "raw.xlsx")
    book = openpyxl.load_workbook(path); book["Key SKUs List"]["C3"] = None; book.save(path)
    adapter = WorkbookAdapter(path, CONTRACT)
    try:
        mapping, resolutions = adapter.resolve_mapping("JD", "S3203/08")
        assert not any(item.brand == "FLYCO" for item in mapping)
        assert next(item for item in resolutions if item.brand == "FLYCO").state == "NO_MAPPING_DECLARED"
    finally:
        adapter.close()


def test_multiple_models_require_confirmation(tmp_path: Path) -> None:
    path = make_fixture(tmp_path / "raw.xlsx")
    book = openpyxl.load_workbook(path); key = book["Key SKUs List"]; key["C3"] = "FS966/967/968"; sheet = book["3.2 Shaver SKU(JD)"]
    rows = [row for row in range(1, sheet.max_row + 1) if sheet.cell(row, 13).value == "FS966/967/968"]
    for row in reversed(rows):
        sheet.cell(row, 13, "FS966"); sheet.insert_rows(row + 1)
        for col in range(1, 22): sheet.cell(row + 1, col, sheet.cell(row, col).value)
        sheet.cell(row + 1, 13, "FS967")
    book.save(path)
    adapter = WorkbookAdapter(path, CONTRACT)
    try:
        mapping, resolutions = adapter.resolve_mapping("JD", "S3203/08")
        flyco = next(item for item in resolutions if item.brand == "FLYCO")
        assert flyco.state == "MULTIPLE_CANDIDATES"
        assert set(flyco.present_candidates) >= {"FS966", "FS967"}
        assert not any(item.brand == "FLYCO" for item in mapping)
    finally:
        adapter.close()


def test_mapping_continuation_rows_are_part_of_the_same_philips_sku(tmp_path: Path) -> None:
    path = make_fixture(tmp_path / "raw.xlsx")
    book = openpyxl.load_workbook(path)
    key = book["Key SKUs List"]
    key.insert_rows(4)
    key["C3"] = "FS966"
    key["C4"] = "FS967"
    sheet = book["3.2 Shaver SKU(JD)"]
    for row in range(1, sheet.max_row + 1):
        if sheet.cell(row, 13).value == "FS966/967/968":
            sheet.cell(row, 13, "FS966")
            sheet.insert_rows(row + 1)
            for col in range(1, 22):
                sheet.cell(row + 1, col, sheet.cell(row, col).value)
            sheet.cell(row + 1, 13, "FS967")
    book.save(path)
    adapter = WorkbookAdapter(path, CONTRACT)
    try:
        mappings, resolutions = adapter.resolve_mapping("JD", "S3203/08")
        resolution = next(item for item in resolutions if item.brand == "FLYCO")
        assert resolution.state == "MULTIPLE_CANDIDATES"
        assert resolution.present_candidates == ["FS966", "FS967"]
        assert not any(item.brand == "FLYCO" for item in mappings)
    finally:
        adapter.close()


def test_identical_duplicate_competitor_rows_are_collapsed_with_provenance(tmp_path: Path) -> None:
    path = make_fixture(tmp_path / "raw.xlsx")
    book = openpyxl.load_workbook(path)
    sheet = book["3.2 Shaver SKU(JD)"]
    source_row = next(row for row in range(1, sheet.max_row + 1) if sheet.cell(row, 13).value == "5603")
    sheet.insert_rows(source_row + 1)
    for col in range(1, 22):
        sheet.cell(source_row + 1, col, sheet.cell(source_row, col).value)
    book.save(path)
    adapter = WorkbookAdapter(path, CONTRACT)
    try:
        mappings, _ = adapter.resolve_mapping("JD", "S3203/08")
        records = adapter.extract("JD", mappings, 12)
        braun = next(item for item in records if item.brand == "BRAUN")
        assert braun.source_cells["price"].duplicate_cells
    finally:
        adapter.close()


@given(prefix=st.from_regex(r"[A-Z]{1,4}", fullmatch=True), first=st.integers(100, 999), rest=st.lists(st.integers(100, 999), min_size=1, max_size=3, unique=True))
def test_prefix_aware_slash_expansion(prefix: str, first: int, rest: list[int]) -> None:
    raw = "/".join([f"{prefix}{first}", *map(str, rest)])
    assert WorkbookAdapter._expand_slash_models(raw) == [f"{prefix}{first}", *[f"{prefix}{item}" for item in rest]]


def test_claims_are_evidence_bounded(tmp_path: Path) -> None:
    path = make_fixture(tmp_path / "raw.xlsx"); adapter = WorkbookAdapter(path, CONTRACT)
    try:
        mapping, _ = adapter.resolve_mapping("JD", "S3203/08"); records = adapter.extract("JD", mapping, 12); claims = deterministic_claims(build_observations(records, mapping), [])
        assert len(claims) == 10 and claims[-1].claim_type == "decision"
        assert claims[6].evidence_level == "HYPOTHESIS" and not claims[6].external_evidence_ids
    finally:
        adapter.close()


def test_duckdb_keeps_provenance_and_claims(tmp_path: Path) -> None:
    path = make_fixture(tmp_path / "raw.xlsx"); adapter = WorkbookAdapter(path, CONTRACT)
    try:
        mapping, _ = adapter.resolve_mapping("JD", "S1115/02"); records = adapter.extract("JD", mapping, 12); claims = deterministic_claims(build_observations(records, mapping), []); database = tmp_path / "analysis.duckdb"; quality = validate_metrics(records, "S1115/02", 12, adapter.config.data["quality"])
        facts = build_standard_facts(records, "JD"); DuckDBRepository(database).save("run-1", "JD", "S1115/02", "hash", mapping, records, [], claims, quality, facts)
        connection = duckdb.connect(str(database), read_only=True)
        assert connection.execute("select count(*) from monthly_metric").fetchone()[0] == 60
        assert connection.execute("select model from sku_mapping where brand='FLYCO'").fetchone()[0] == "FS903"
        assert connection.execute("select source_cells_json from monthly_metric limit 1").fetchone()[0]
        assert connection.execute("select count(*) from insight_claim").fetchone()[0] == 10
        assert connection.execute("select count(*) from standard_fact").fetchone()[0] == 60
        connection.close()
    finally:
        adapter.close()


def test_unknown_structure_fails_closed(tmp_path: Path) -> None:
    path = make_fixture(tmp_path / "raw.xlsx"); book = openpyxl.load_workbook(path); del book["3.1 Shaver SKU(ALi)"]; book.save(path)
    with pytest.raises(WorkbookContractError) as error: WorkbookAdapter(path, CONTRACT)
    assert error.value.code == "UNSUPPORTED_STRUCTURE"


def test_horizontally_aligned_channel_skus_are_not_aliases(tmp_path: Path) -> None:
    path = make_fixture(tmp_path / "raw.xlsx")
    book = openpyxl.load_workbook(path)
    key = book["Key SKUs List"]
    key.delete_rows(1, key.max_row)
    groups = {"ALI": 2, "JD": 8, "OFFLINE": 14}
    for channel, start_col in groups.items():
        key.cell(3, start_col, channel)
        for offset, brand in enumerate(("PHILIPS", "BRAUN", "FLYCO", "PANASONIC", "YOOSE")):
            key.cell(4, start_col + offset, brand)
    key["B5"] = "S1113"
    key["D5"] = "FS903"
    key["H5"] = "S1115/02"
    key["J5"] = "FS903"
    key["N5"] = "S1213/02"
    ali = book["3.1 Shaver SKU(ALi)"]
    offline = book["3.3 Shaver SKU(Offline)"]
    for row in range(1, ali.max_row + 1):
        if ali.cell(row, 2).value == "S1115/02":
            ali.cell(row, 2, "S1113")
    for row in range(1, offline.max_row + 1):
        if offline.cell(row, 2).value == "S1115/02":
            offline.cell(row, 2, "S1213/02")
    book.save(path)

    adapter = WorkbookAdapter(path, CONTRACT)
    try:
        for channel, sku in (("ALI", "S1113"), ("JD", "S1115/02"), ("OFFLINE", "S1213/02")):
            mappings, _ = adapter.resolve_mapping(channel, sku)
            primary = next(item for item in mappings if item.role == "PH")
            assert primary.model == sku
            assert primary.mapping_state == "EXACT_SINGLE"
            assert primary.source == "Key SKUs List"

        with pytest.raises(WorkbookContractError) as ali_error:
            adapter.resolve_mapping("ALI", "S1115/02")
        assert ali_error.value.code == "PRIMARY_SKU_NOT_FOUND_IN_CHANNEL"

        with pytest.raises(WorkbookContractError) as offline_error:
            adapter.resolve_mapping("OFFLINE", "S1115/02")
        assert offline_error.value.code == "PRIMARY_SKU_NOT_FOUND_IN_CHANNEL"
    finally:
        adapter.close()
