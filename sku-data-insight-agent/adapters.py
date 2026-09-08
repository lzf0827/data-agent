from __future__ import annotations

import json
import math
import re
import zipfile
from dataclasses import asdict, dataclass, field
from datetime import date
from pathlib import Path
from typing import Any

import openpyxl
import yaml

from semantic import WorkbookSemanticPlan, WorkbookSemanticProfiler


MONTHS = {name: index + 1 for index, name in enumerate(("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"))}
PRIMARY_MODEL_PATTERN = re.compile(r"(?<![A-Z0-9])([A-Z]{1,5}\d{2,6})(?:/(\d{2}))?", re.I)


def normalize(value: Any) -> str:
    return re.sub(r"[^A-Z0-9]", "", str(value or "").upper())


def primary_model_tokens(value: Any) -> set[str]:
    """Canonicalize Philips product labels without treating series descriptors as SKUs."""
    text = str(value or "").upper()
    text = re.sub(r"\b([A-Z]{1,5}\d{2,6})-(\d{2})\b", r"\1/\2", text)
    matches = list(PRIMARY_MODEL_PATTERN.finditer(text))
    if len(matches) > 1:
        specific = [match for match in matches if not re.fullmatch(r"[SI]\d000", match.group(1), re.I)]
        matches = specific or matches

    tokens: set[str] = set()
    component: list[re.Match[str]] = []

    def flush_component() -> None:
        if not component:
            return
        suffixes = [match.group(2) for match in component if match.group(2)]
        shared_suffix = suffixes[-1] if len(component) > 1 and any(not match.group(2) for match in component) else None
        for match in component:
            suffix = match.group(2) or shared_suffix
            tokens.add(normalize(f"{match.group(1)}/{suffix}" if suffix else match.group(1)))

    for match in matches:
        if component:
            separator = text[component[-1].end():match.start()]
            if not re.fullmatch(r"\s*/\s*", separator):
                flush_component()
                component = []
        component.append(match)
    flush_component()

    if matches:
        first_prefix = re.match(r"[A-Z]+", matches[0].group(1), re.I)
        if first_prefix:
            for shorthand in re.finditer(r"(?<![A-Z0-9])/\s*(\d{3,6})/(\d{2})(?!\d)", text):
                tokens.add(normalize(f"{first_prefix.group(0)}{shorthand.group(1)}/{shorthand.group(2)}"))
    return tokens


def primary_token_set_matches(requested: set[str], candidate: set[str]) -> bool:
    """Match explicit suffixes strictly; a suffixless single label may match its unique family member."""
    for token in requested:
        if token in candidate:
            continue
        if any(other.startswith(token) and len(other) == len(token) + 2 for other in candidate):
            continue
        return False
    return True


def clean_number(value: Any) -> float | None:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    numeric = float(value)
    return numeric if math.isfinite(numeric) else None


def model_matches(raw_value: Any, expected: str) -> bool:
    def aliases(value: Any) -> set[str]:
        text = str(value or "").strip()
        if not text:
            return set()
        result = {normalize(text)}
        parenthetical_base = re.sub(r"\s*\([^)]*\)\s*$", "", text).strip()
        if parenthetical_base and parenthetical_base != text:
            result.add(normalize(parenthetical_base))
        if "/" in text:
            parts = [part.strip() for part in text.split("/") if part.strip()]
            first_match = re.match(r"^(.*?)(\d+)$", parts[0]) if parts else None
            if first_match and len(parts) > 1 and all(re.fullmatch(r"\d+", part) for part in parts[1:]):
                prefix = first_match.group(1)
                result.update(normalize(item) for item in [parts[0], *[f"{prefix}{part}" for part in parts[1:]]])
            elif len(parts) > 1 and all(re.search(r"[A-Z]", part, re.I) for part in parts):
                result.update(normalize(part) for part in parts)
        return {item for item in result if item}

    raw_aliases = aliases(raw_value)
    expected_aliases = aliases(expected)
    return bool(raw_aliases & expected_aliases)


@dataclass(frozen=True)
class MappingItem:
    brand: str
    model: str
    role: str
    source: str
    confidence: float = 1.0
    mapping_state: str = "EXACT_SINGLE"
    raw_mapping_value: str | None = None
    source_cells: tuple[str, ...] = ()


@dataclass(frozen=True)
class MappingResolution:
    brand: str
    raw_value: str | None
    candidates: list[str]
    present_candidates: list[str]
    state: str
    selected_model: str | None
    reason: str
    declared_cells: list[str] = field(default_factory=list)
    available_periods: dict[str, list[str]] = field(default_factory=dict)
    coverage: dict[str, float] = field(default_factory=dict)


@dataclass(frozen=True)
class CellProvenance:
    sheet: str
    cell: str
    raw_value: float | None
    formula: str | None
    duplicate_cells: tuple[str, ...] = ()


@dataclass(frozen=True)
class MonthlyMetric:
    period: date
    period_label: str
    brand: str
    model: str
    role: str
    price: float | None
    sales: float | None
    sales_value: float | None
    source_sheet: str
    source_cells: dict[str, CellProvenance]
    mapping_state: str

    def to_dict(self) -> dict[str, Any]:
        value = asdict(self)
        value["period"] = self.period.isoformat()
        return value


class WorkbookContractError(ValueError):
    def __init__(self, code: str, message: str, details: dict[str, Any] | None = None) -> None:
        super().__init__(message)
        self.code = code
        self.details = details or {}


class ContractConfig:
    def __init__(self, path: Path) -> None:
        self.path = path
        self.data = yaml.safe_load(path.read_text(encoding="utf-8"))

    @property
    def adapter_version(self) -> str:
        return str(self.data["adapter_version"])

    def channel_sheet(self, channel: str) -> str:
        try:
            return str(self.data["channels"][channel]["sheet"])
        except KeyError as exc:
            raise WorkbookContractError("UNSUPPORTED_CHANNEL", f"Unsupported channel: {channel}") from exc


class WorkbookAdapter:
    def __init__(self, path: Path, contract_path: Path) -> None:
        self.path = path
        self.config = ContractConfig(contract_path)
        self._preflight_container()
        self.values_book = openpyxl.load_workbook(path, read_only=False, data_only=True, keep_links=False)
        self.formula_book = openpyxl.load_workbook(path, read_only=False, data_only=False, keep_links=False)
        self.detect_version()

    def close(self) -> None:
        self.values_book.close()
        self.formula_book.close()

    def _preflight_container(self) -> None:
        limit = int(self.config.data["workbook"]["maximum_file_bytes"])
        if self.path.suffix.lower() != ".xlsx" or not self.path.exists():
            raise WorkbookContractError("INVALID_FILE_TYPE", "Only an existing .xlsx workbook is supported")
        if self.path.stat().st_size > limit:
            raise WorkbookContractError("FILE_TOO_LARGE", f"Workbook exceeds {limit} bytes")
        try:
            with zipfile.ZipFile(self.path) as archive:
                total = sum(item.file_size for item in archive.infolist())
                if total > limit * 20:
                    raise WorkbookContractError("UNSAFE_XLSX_CONTAINER", "Expanded workbook is unexpectedly large")
                lowered = {item.filename.lower() for item in archive.infolist()}
                if any("vbaproject.bin" in name for name in lowered):
                    raise WorkbookContractError("MACRO_NOT_ALLOWED", "Embedded VBA is not allowed")
        except PermissionError as exc:
            raise WorkbookContractError("SOURCE_FILE_LOCKED", "The workbook cannot be read. Close it in Excel or upload a copy, then retry.") from exc
        except zipfile.BadZipFile as exc:
            raise WorkbookContractError("INVALID_XLSX_CONTAINER", "Workbook is not a valid XLSX container") from exc

    def detect_version(self) -> str:
        required = set(self.config.data["workbook"]["required_sheets"])
        missing = sorted(required - set(self.values_book.sheetnames))
        if missing:
            raise WorkbookContractError("UNSUPPORTED_STRUCTURE", "Required sheets are missing", {"missing_sheets": missing})
        key_sheet = self.values_book[self.config.data["mapping"]["sheet"]]
        if not any(cell.value not in (None, "") for row in key_sheet.iter_rows() for cell in row):
            raise WorkbookContractError("UNSUPPORTED_STRUCTURE", "Key SKUs List is empty")
        minimum = int(self.config.data["workbook"]["minimum_month_blocks"])
        for channel in self.config.data["channels"]:
            sheet = self.values_book[self.config.channel_sheet(channel)]
            if len(self._month_blocks(sheet)) < minimum:
                raise WorkbookContractError("UNSUPPORTED_STRUCTURE", f"{channel} has fewer than {minimum} month blocks")
        return self.config.adapter_version

    def _month_blocks(self, sheet: Any) -> list[tuple[int, date, str]]:
        header_col = int(self.config.data["layout"]["month_header_column"])
        blocks: list[tuple[int, date, str]] = []
        for row in range(1, sheet.max_row + 1):
            header = str(sheet.cell(row, header_col).value or "").strip()
            match = re.fullmatch(r"([A-Z][a-z]{2})(\d{4}) Value", header)
            if match and match.group(1) in MONTHS:
                period = date(int(match.group(2)), MONTHS[match.group(1)], 1)
                blocks.append((row, period, f"{str(period.year)[2:]}-{period.month:02d}"))
        return blocks

    def model_inventory(self, channel: str) -> dict[str, list[str]]:
        sheet = self.values_book[self.config.channel_sheet(channel)]
        layout = self.config.data["layout"]["competitor"]
        result: dict[str, list[str]] = {}
        for row in range(1, sheet.max_row + 1):
            model = sheet.cell(row, int(layout["model"])).value
            brand = sheet.cell(row, int(layout["brand"])).value
            if model in (None, "") or brand in (None, ""):
                continue
            brand_key = str(brand).strip().upper()
            model_text = str(model).strip()
            if model_text not in result.setdefault(brand_key, []):
                result[brand_key].append(model_text)
        return result

    def primary_model_inventory(self, channel: str) -> list[str]:
        sheet = self.values_book[self.config.channel_sheet(channel)]
        model_col = int(self.config.data["layout"]["primary"]["model"])
        models: list[str] = []
        for row in range(1, sheet.max_row + 1):
            value = sheet.cell(row, model_col).value
            if value in (None, ""):
                continue
            model = str(value).strip()
            if model and not any(normalize(item) == normalize(model) for item in models):
                models.append(model)
        return models

    def semantic_profiler(self) -> WorkbookSemanticProfiler:
        semantic = self.config.data.get("semantic", {})
        aliases = semantic.get("channel_aliases", {channel: [channel] for channel in self.config.data["channels"]})
        primary_aliases = semantic.get("primary_header_aliases", ["PHILIPS"])
        return WorkbookSemanticProfiler(
            self.values_book,
            self.config.data["mapping"]["sheet"],
            {channel: self.config.channel_sheet(channel) for channel in self.config.data["channels"]},
            aliases,
            primary_aliases,
            float(semantic.get("label_confidence_threshold", 0.90)),
        )

    @staticmethod
    def _expand_slash_models(raw: str) -> list[str]:
        raw = raw.strip()
        if "/" not in raw:
            return [raw] if raw else []
        parts = [part.strip() for part in raw.split("/") if part.strip()]
        if len(parts) < 2:
            return [raw]
        first = parts[0]
        prefix_match = re.match(r"^(.*?)(\d+)$", first)
        if not prefix_match or not all(re.fullmatch(r"\d+", part) for part in parts[1:]):
            return parts
        prefix = prefix_match.group(1)
        return [first, *[f"{prefix}{part}" for part in parts[1:]]]

    def parse_mapping_candidates(self, raw_value: Any, inventory: list[str]) -> tuple[list[str], list[str]]:
        raw = str(raw_value or "").strip()
        if not raw:
            return [], []
        if any(normalize(item) == normalize(raw) for item in inventory):
            return [raw], [next(item for item in inventory if normalize(item) == normalize(raw))]
        delimiters = self.config.data["mapping"]["competitor_delimiters"]
        pattern = "|".join(re.escape(item) for item in delimiters)
        chunks = [part.strip() for part in re.split(pattern, raw) if part.strip()]
        candidates: list[str] = []
        for chunk in chunks:
            candidates.extend(self._expand_slash_models(chunk))
        candidates = list(dict.fromkeys(candidates))
        present: list[str] = []
        for candidate in candidates:
            exact_matches = [item for item in inventory if normalize(item) == normalize(candidate)]
            candidate_inventory = exact_matches or [item for item in inventory if model_matches(item, candidate)]
            for item in candidate_inventory:
                if model_matches(item, candidate) and item not in present:
                    present.append(item)
        return candidates, present

    def _primary_expression_match(self, raw_value: Any, requested_sku: str, inventory: list[str]) -> str | None:
        exact_inventory = next((item for item in inventory if normalize(item) == normalize(requested_sku)), None)
        if exact_inventory is None:
            return None
        raw = str(raw_value or "").strip()
        if normalize(raw) == normalize(requested_sku):
            return exact_inventory
        if "/" in raw and normalize(requested_sku) in normalize(raw):
            return exact_inventory
        return None

    @staticmethod
    def _primary_tokens(value: Any) -> set[str]:
        return primary_model_tokens(value)

    def primary_model_period_availability(self, channel: str, model: str, months: int) -> list[str]:
        sheet = self.values_book[self.config.channel_sheet(channel)]
        model_col = int(self.config.data["layout"]["primary"]["model"])
        all_blocks = self._month_blocks(sheet)
        blocks = sorted(all_blocks, key=lambda item: item[1])[-months:]
        available: list[str] = []
        for start_row, _, period_label in blocks:
            next_row = min((item[0] for item in all_blocks if item[0] > start_row), default=sheet.max_row + 1)
            if any(normalize(sheet.cell(row, model_col).value) == normalize(model) for row in range(start_row + 1, next_row)):
                available.append(period_label)
        return available

    def _resolve_primary_product_group(self, channel: str, requested_sku: str, inventory: list[str], months: int) -> str | None:
        requested_tokens = self._primary_tokens(requested_sku)
        if not requested_tokens:
            return None
        candidates = [item for item in inventory if primary_token_set_matches(requested_tokens, self._primary_tokens(item))]
        if len(candidates) == 1:
            return candidates[0]
        if len(candidates) > 1:
            availability = {item: self.primary_model_period_availability(channel, item, months) for item in candidates}
            best_coverage = max(len(periods) for periods in availability.values())
            best = [item for item in candidates if len(availability[item]) == best_coverage]
            if best_coverage > 0 and len(best) == 1:
                return best[0]
            exact_token_groups = [item for item in best if self._primary_tokens(item) == requested_tokens]
            if len(exact_token_groups) == 1 and all(len(availability[item]) == 0 for item in best):
                return exact_token_groups[0]
            raise WorkbookContractError(
                "AMBIGUOUS_PRIMARY_PRODUCT_GROUP",
                f"Philips product expression {requested_sku} matches more than one Raw Data product group",
                {
                    "channel": channel,
                    "requested_sku": requested_sku,
                    "candidate_groups": candidates,
                    "requested_period_availability": availability,
                },
            )
        return None

    def _mapping_block_end(self, group: Any, row: int) -> int:
        for merged in self.values_book[self.config.data["mapping"]["sheet"]].merged_cells.ranges:
            if merged.min_col == group.start_column and merged.min_row <= row <= merged.max_row:
                return min(group.data_end_row + 1, merged.max_row + 1)
        sheet = self.values_book[self.config.data["mapping"]["sheet"]]
        end = row + 1
        while end <= group.data_end_row and sheet.cell(end, group.start_column).value in (None, ""):
            end += 1
        return end

    def model_period_availability(self, channel: str, brand: str, model: str, months: int) -> list[str]:
        sheet = self.values_book[self.config.channel_sheet(channel)]
        layout = self.config.data["layout"]["competitor"]
        blocks = sorted(self._month_blocks(sheet), key=lambda item: item[1])[-months:]
        available: list[str] = []
        for start_row, _, period_label in blocks:
            next_row = min((item[0] for item in self._month_blocks(sheet) if item[0] > start_row), default=sheet.max_row + 1)
            if any(
                model_matches(sheet.cell(row, int(layout["model"])).value, model)
                and normalize(sheet.cell(row, int(layout["brand"])).value) == normalize(brand)
                for row in range(start_row + 1, next_row)
            ):
                available.append(period_label)
        return available

    def resolve_mapping(
        self,
        channel: str,
        sku: str,
        approved: dict[str, str] | None = None,
        *,
        semantic_plan: WorkbookSemanticPlan | None = None,
        months: int = 12,
    ) -> tuple[list[MappingItem], list[MappingResolution]]:
        approved = {str(key).upper(): str(value) for key, value in (approved or {}).items()}
        sheet = self.values_book[self.config.data["mapping"]["sheet"]]
        profiler = self.semantic_profiler()
        if semantic_plan is None:
            manifest = profiler.build_manifest()
            semantic_plan = profiler.validate_plan(profiler.deterministic_plan(manifest), manifest)
        try:
            group = semantic_plan.group_for(channel)
        except KeyError as exc:
            raise WorkbookContractError("MAPPING_GROUP_NOT_FOUND", f"Mapping group not found: {channel}") from exc
        primary_inventory = self.primary_model_inventory(channel)
        primary_model = self._resolve_primary_product_group(channel, sku, primary_inventory, months)
        product_group_tokens = self._primary_tokens(primary_model) if primary_model else set()
        target_rows: list[tuple[int, str]] = []
        if primary_model:
            for row in range(group.data_start_row, group.data_end_row + 1):
                mapping_tokens = self._primary_tokens(sheet.cell(row, group.start_column).value)
                if mapping_tokens and primary_token_set_matches(mapping_tokens, product_group_tokens):
                    target_rows.append((row, primary_model))
        distinct_mapping_keys = {
            normalize(sheet.cell(row, group.start_column).value)
            for row, _ in target_rows
            if sheet.cell(row, group.start_column).value not in (None, "")
        }
        if len(distinct_mapping_keys) > 1:
            raise WorkbookContractError(
                "AMBIGUOUS_PRODUCT_GROUP_MAPPING",
                f"Philips product group {primary_model} has more than one mapping representative in {channel}",
                {
                    "channel": channel,
                    "requested_sku": sku,
                    "raw_product_group": primary_model,
                    "mapping_cells": [sheet.cell(row, group.start_column).coordinate for row, _ in target_rows],
                },
            )
        if not target_rows:
            raise WorkbookContractError(
                "PRIMARY_SKU_NOT_FOUND_IN_CHANNEL",
                f"Philips SKU {sku} is not declared in the {channel} mapping group",
                {
                    "channel": channel,
                    "sku": sku,
                    "raw_product_group": primary_model,
                    "mapping_sheet": self.config.data["mapping"]["sheet"],
                    "source_sheet": self.config.channel_sheet(channel),
                    "cross_channel_fallback": False,
                    "semantic_plan_fingerprint": semantic_plan.layout_fingerprint,
                },
            )
        inventory = self.model_inventory(channel)
        primary_cells = tuple(sheet.cell(row, group.start_column).coordinate for row, _ in target_rows)
        raw_primary_values = "\n".join(dict.fromkeys(str(sheet.cell(row, group.start_column).value).strip() for row, _ in target_rows))
        primary_state = "EXACT_SINGLE" if normalize(primary_model) == normalize(sku) and normalize(raw_primary_values) == normalize(sku) else "PRODUCT_GROUP_ALIAS"
        primary_source = "Key SKUs List" if primary_state == "EXACT_SINGLE" else "Key SKUs List + channel Raw Data"
        mappings = [MappingItem("PHILIPS", primary_model, "PH", primary_source, 1.0, primary_state, raw_primary_values, primary_cells)]
        resolutions: list[MappingResolution] = []
        requested_period_count = min(months, len(self._month_blocks(self.values_book[self.config.channel_sheet(channel)])))
        minimum_coverage = float(self.config.data["quality"].get("competitor_minimum_coverage", 0.75))
        for header in group.competitor_headers:
            brand = header.brand
            column = self.values_book[self.config.data["mapping"]["sheet"]][header.label.cell].column
            declared_cells: list[str] = []
            raw_values: list[str] = []
            for target_row, _ in target_rows:
                for row in range(target_row, self._mapping_block_end(group, target_row)):
                    value = sheet.cell(row, column).value
                    if value in (None, ""):
                        continue
                    raw_values.append(str(value).strip())
                    declared_cells.append(sheet.cell(row, column).coordinate)
            raw_text = "\n".join(dict.fromkeys(raw_values)) or None
            candidates, present = self.parse_mapping_candidates(raw_text, inventory.get(brand, []))
            availability = {model: self.model_period_availability(channel, brand, model, months) for model in present}
            coverage = {model: len(periods) / requested_period_count if requested_period_count else 0.0 for model, periods in availability.items()}
            viable = [model for model in present if coverage[model] >= minimum_coverage]
            selected = None
            state = "NO_MAPPING_DECLARED"
            reason = "Key SKUs List cell is empty"
            approved_model = approved.get(brand)
            approved_match = next((item for item in present if approved_model and model_matches(item, approved_model)), None)
            if approved_match:
                selected, state, reason = approved_match, "CONFIRMED_OVERRIDE", "Employee approved one declared physical model"
            elif raw_text is not None and not present:
                state, reason = "RAW_MODEL_NOT_FOUND", "Declared model is absent from the selected channel Raw Data"
            elif len(viable) == 1:
                selected, state, reason = viable[0], "EXACT_SINGLE", "One declared model meets requested-period coverage"
            elif len(viable) > 1:
                state, reason = "MULTIPLE_CANDIDATES", "More than one declared model meets coverage and requires confirmation"
            elif present:
                union_periods = set().union(*(availability[model] for model in present))
                if len(present) > 1 and requested_period_count and len(union_periods) / requested_period_count >= minimum_coverage:
                    state, reason = "LIFECYCLE_SPLIT_REQUIRED", "No single model covers the period; declared models collectively span it"
                else:
                    state, reason = "MODEL_NOT_AVAILABLE_FOR_PERIOD", "Declared model coverage is below the required threshold"
            resolution = MappingResolution(brand, raw_text, candidates, present, state, selected, reason, declared_cells, availability, coverage)
            resolutions.append(resolution)
            if selected:
                mappings.append(MappingItem(brand, selected, "Competitor", "Key SKUs List" if state == "EXACT_SINGLE" else "confirmed override", 1.0, state, raw_text, tuple(declared_cells)))
        return mappings, resolutions

    def extract(self, channel: str, mappings: list[MappingItem], months: int) -> list[MonthlyMetric]:
        sheet_name = self.config.channel_sheet(channel)
        sheet = self.values_book[sheet_name]
        formula_sheet = self.formula_book[sheet_name]
        blocks = sorted(self._month_blocks(sheet), key=lambda item: item[1])
        if len(blocks) < months:
            raise WorkbookContractError("INSUFFICIENT_MONTHS", f"Requested {months} months but found {len(blocks)}")
        selected = blocks[-months:]
        if len({item[1] for item in selected}) != months:
            raise WorkbookContractError("DUPLICATE_MONTH", "Selected month blocks are not unique")
        layout = self.config.data["layout"]
        records: list[MonthlyMetric] = []
        for block_position, (start_row, period, period_label) in enumerate(selected):
            # The production workbook stores newest months first. Date order chooses
            # the requested periods, while physical row order owns block boundaries.
            next_row = min((item[0] for item in blocks if item[0] > start_row), default=sheet.max_row + 1)
            for mapping in mappings:
                section = layout["primary"] if mapping.role == "PH" else layout["competitor"]
                matches: list[int] = []
                for row in range(start_row + 1, next_row):
                    model_value = sheet.cell(row, int(section["model"])).value
                    if mapping.role == "PH":
                        matched = normalize(model_value) == normalize(mapping.model)
                    else:
                        brand_value = sheet.cell(row, int(section["brand"])).value
                        matched = model_matches(model_value, mapping.model) and normalize(brand_value) == normalize(mapping.brand)
                    if matched:
                        matches.append(row)
                if len(matches) > 1:
                    metric_columns = [int(section[key]) for key in ("value", "unit", "asp")]
                    signatures = {
                        tuple(clean_number(sheet.cell(candidate_row, col).value) for col in metric_columns)
                        for candidate_row in matches
                    }
                    if len(signatures) > 1:
                        raise WorkbookContractError(
                            "CONFLICTING_DUPLICATE_MODEL_ROW",
                            f"{mapping.brand} {mapping.model} has conflicting duplicate rows in {period_label}",
                            {"sheet": sheet_name, "rows": matches, "signatures": [list(item) for item in signatures]},
                        )
                row = matches[0] if matches else None
                source_cells: dict[str, CellProvenance] = {}
                metric_values: dict[str, float | None] = {}
                for metric, key in (("sales_value", "value"), ("sales", "unit"), ("price", "asp")):
                    if row is None:
                        source_cells[metric] = CellProvenance(sheet_name, "", None, None)
                        metric_values[metric] = None
                        continue
                    col = int(section[key])
                    value_cell = sheet.cell(row, col)
                    formula_cell = formula_sheet.cell(row, col)
                    formula = str(formula_cell.value) if isinstance(formula_cell.value, str) and formula_cell.value.startswith("=") else None
                    numeric = clean_number(value_cell.value)
                    if formula and numeric is None:
                        raise WorkbookContractError("FORMULA_CACHE_MISSING", f"Formula cell {sheet_name}!{value_cell.coordinate} has no cached numeric value")
                    duplicate_cells = tuple(sheet.cell(other_row, col).coordinate for other_row in matches[1:])
                    source_cells[metric] = CellProvenance(sheet_name, value_cell.coordinate, numeric, formula, duplicate_cells)
                    metric_values[metric] = numeric
                records.append(MonthlyMetric(period, period_label, mapping.brand, mapping.model, mapping.role, metric_values["price"], metric_values["sales"], metric_values["sales_value"], sheet_name, source_cells, mapping.mapping_state))
        return records

def serialize_mappings(mappings: list[MappingItem]) -> list[dict[str, Any]]:
    return [asdict(item) for item in mappings]


def serialize_resolutions(resolutions: list[MappingResolution]) -> list[dict[str, Any]]:
    return [asdict(item) for item in resolutions]


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
